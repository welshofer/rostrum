"""Absolute cost and independent saved-grid validation of the new rotated path.
Run from the repository root: python3 THIS.py SCRATCH OUTPUT.json.
Build SCRATCH/rotated from the sibling Swift driver and SCRATCH/measure-child first.
No concurrent builds/tests. These timings are not a before/after speedup claim.
"""
import hashlib,json,os,platform,statistics,subprocess,sys,tempfile,time
from pathlib import Path
from pptx import Presentation
scratch,output=map(Path,sys.argv[1:3])
if hasattr(os,'sched_setaffinity'):os.sched_setaffinity(0,{min(os.sched_getaffinity(0))})
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
r={'platform':platform.platform(),'binary_sha256':sha(scratch/'rotated'),
   'driver_sha256':sha(Path('docs/benchmarks/2026-10-17-rotated-tables.swift')),
   'cpu_affinity':sorted(os.sched_getaffinity(0)) if hasattr(os,'sched_getaffinity') else None,
   'preservation':[],'runs':[]}
for rows,labels in [(24,'distinct'),(28,'mixed'),(100,'distinct')]:
 for sizing in ['estimated','measured']:
  paths=[]
  for repeat in range(2):
   p=scratch/f'rotated-{rows}-{labels}-{sizing}-{repeat}'
   assert not p.exists()
   subprocess.run([str(scratch/'rotated'),str(rows),labels,sizing,'1','artifacts',str(p),'--rotated'],check=True,stdout=subprocess.DEVNULL)
   paths.append(p)
  names=sorted(p.name for p in paths[0].iterdir())
  assert names==sorted(p.name for p in paths[1].iterdir())
  for name in names:
   a,b=[(p/name).read_bytes() for p in paths]
   assert json.loads(a)==json.loads(b) if name.endswith('.json') else a==b,name
  p=paths[0];deck=Presentation(p/'saved.pptx');source=json.loads((p/'source.json').read_text())
  directions={(d['row'],d['column']):d['direction'] for d in json.loads((p/'directions.json').read_text())}
  pages=json.loads((p/'geometry.json').read_text())
  assert [i for page in pages for i in page['body_rows']]==list(range(2,rows+2))
  for index,page in enumerate(pages):
   shape,=[s for s in deck.slides[index+1].shapes if s.has_table];table=shape.table
   assert shape.top+shape.height<=deck.slide_height
   assert [row.height for row in table.rows]==page['heights']
   assert [col.width for col in table.columns]==page['widths']
   for local,original in enumerate(page['source_rows']):
    for column,cell in enumerate(table.rows[local].cells):
     assert cell.text==source[original][column]
     assert cell._tc.get_or_add_tcPr().get('vert','horz')==directions.get((original,column),'horz')
   for cell in page['cells']:
    actual=table.cell(cell['row'],cell['column'])
    assert actual.is_spanned==cell['covered']
    if not cell['covered']:
     assert actual.span_height==cell['rowSpan'] and actual.span_width==cell['span']
     assert sum(page['heights'][cell['row']:cell['row']+cell['rowSpan']])>=cell['height']
  r['preservation'].append({'rows':rows,'labels':labels,'sizing':sizing,'pages':len(pages),'identical_files':len(names),'independent_grid_direction_source_checks':True})
  print('Rotated saved-grid checks:',rows,labels,sizing,flush=True)
for batch in range(3):
 for rows in [24,1000]:
  for sizing in (['estimated','measured'] if batch%2==0 else ['measured','estimated']):
   with tempfile.NamedTemporaryFile() as usage:
    start=time.perf_counter()
    p=subprocess.run([str(scratch/'measure-child'),usage.name,str(scratch/'rotated'),str(rows),'distinct',sizing,'5','compose','--rotated'],capture_output=True,text=True,check=True)
    r['runs'].append({'batch':batch,'rows':rows,'sizing':sizing,'samples':json.loads(p.stdout),
                     'process_elapsed_ms':(time.perf_counter()-start)*1000,**json.load(usage)})
   output.write_text(json.dumps(r,indent=2,sort_keys=True)+'\n')
 print('Rotated timing batch',batch+1,'of 3',flush=True)
for rows in [24,1000]:
 for sizing in ['estimated','measured']:
  records=[x for x in r['runs'] if x['rows']==rows and x['sizing']==sizing]
  print(rows,sizing,{k:statistics.median(x[k] for x in records) for k in ['process_cpu_ms','process_elapsed_ms','rss_kib']})
