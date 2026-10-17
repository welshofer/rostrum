"""Run from the repo root after building the seven drivers documented in the repeated-header report.

python3 docs/benchmarks/2026-10-16-repeated-headers.py SCRATCH OUTPUT.json
Keep builds/tests idle during this alternating fresh-process measurement.
"""
import hashlib
import json
from pathlib import Path
import platform
import statistics
import subprocess
import sys
import tempfile
import time
import zipfile
from pptx import Presentation

scratch, output = Path(sys.argv[1]), Path(sys.argv[2])
launcher = scratch / "measure-child"
subprocess.run(["cc", "-O2", "docs/benchmarks/2026-10-16-measure-child.c", "-o", str(launcher)], check=True)
root = scratch / "preservation"
root.mkdir(exist_ok=False)
record = {"baseline_commit": (scratch / "baseline-commit.txt").read_text().strip(),
          "platform": platform.platform(), "preservation": [], "pairs": [], "headers": []}
record["memory_method"] = "wait4 in a small native launcher; fork after exec avoids inheriting Python peak RSS"
record["swift_version"] = subprocess.check_output(["swift", "--version"], text=True).strip()
record["binary_hashes"] = {name: hashlib.sha256((scratch / name).read_bytes()).hexdigest()
                           for name in ["baseline", "baseline-spans", "candidate", "candidate-spans", "baseline-vertical", "candidate-vertical", "headers"]}
record["source_hashes"] = {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in [
    Path("Sources/Rostrum/Presentation/TableLayoutSource.swift"),
    Path("Sources/RostrumLayout/TableLayoutEngine.swift"), Path("Sources/RostrumLayout/TableCellPreparation.swift"),
    Path("Sources/RostrumLayout/TableRowGroups.swift"), Path("Sources/RostrumLayout/TemplateTableContent.swift"),
    Path("Sources/RostrumLayout/TemplateTablePagination.swift"), Path("Sources/RostrumLayout/TableColumnSizing.swift")]}
record["driver_hashes"] = {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in [
    Path("docs/benchmarks/2026-10-12-table-preparation.swift"), Path("docs/benchmarks/2026-10-12-table-spans.swift"),
    Path("docs/benchmarks/2026-10-16-table-row-pagination.swift"), Path("docs/benchmarks/2026-10-16-repeated-headers.swift"),
    Path(__file__), Path("docs/benchmarks/2026-10-16-measure-child.c")]}

def check_pair(a, b):
    names = sorted(p.name for p in a.iterdir())
    assert names == sorted(p.name for p in b.iterdir())
    for name in names:
        x, y = (a / name).read_bytes(), (b / name).read_bytes()
        assert json.loads(x) == json.loads(y) if name.endswith(".json") else x == y, name
    with zipfile.ZipFile(b / "saved.pptx") as archive:
        assert archive.testzip() is None
    Presentation(b / "saved.pptx")
    return len(names)

# Include fresh font/shaping diagnostics, page boundaries, widths, heights, decks
# and every SVG. Existing JSON keys and ordered arrays must remain identical.
for kind in ["", "-spans", "-vertical"]:
    for rows, labels in [(24, "distinct"), (28, "mixed"), (100, "distinct")]:
        paths = []
        for variant in ["baseline", "candidate"]:
            path = root / f"{variant}{kind}-{rows}-{labels}"
            subprocess.run([str(scratch / (variant + kind)), str(rows), labels, "measured", "1", "artifacts", str(path)],
                           check=True, stdout=subprocess.DEVNULL)
            paths.append(path)
        record["preservation"].append({"kind": kind or "plain", "rows": rows, "labels": labels,
                                       "identical_files": check_pair(*paths)})
for rows in [24, 100]:
    paths = []
    for variant in ["first", "repeat"]:
        path = root / f"headers-{rows}-{variant}"
        subprocess.run([str(scratch / "headers"), str(rows), "distinct", "measured", "1", "artifacts", str(path)],
                       check=True, stdout=subprocess.DEVNULL)
        paths.append(path)
    count = check_pair(*paths)
    deck = Presentation(paths[0] / "saved.pptx")
    pages = json.loads((paths[0] / "geometry.json").read_text())
    assert [row for page in pages for row in page["body_rows"]] == list(range(2, rows + 2))
    source = json.loads((paths[0] / "source.json").read_text())
    for index, geometry in enumerate(pages):
        table, = [s.table for s in deck.slides[index + 1].shapes if s.has_table]
        assert [[c.text for c in r.cells] for r in table.rows] == [source[r] for r in geometry["source_rows"]]
        assert [r.height for r in table.rows] == geometry["heights"]
        assert [c.width for c in table.columns] == geometry["widths"]
        for cell in geometry["cells"]:
            native = table.cell(cell["row"], cell["column"])
            assert native.is_spanned == cell["covered"]
            if not cell["covered"]:
                assert native.span_height == cell["rowSpan"] and native.span_width == cell["span"]
                assert sum(geometry["heights"][cell["row"]:cell["row"] + cell["rowSpan"]]) >= cell["height"]
    record["preservation"].append({"kind": "headers", "rows": rows, "identical_files": count,
                                   "independent_topology_and_geometry": True})

def run(arguments):
    with tempfile.TemporaryFile() as stdout, tempfile.TemporaryFile() as stderr, tempfile.NamedTemporaryFile() as usage:
        start = time.perf_counter()
        p = subprocess.run([str(launcher), usage.name] + [str(a) for a in arguments], stdout=stdout, stderr=stderr)
        elapsed = (time.perf_counter() - start) * 1000
        stderr.seek(0)
        assert p.returncode == 0, stderr.read().decode()
        stdout.seek(0)
        return {"samples": json.load(stdout), "process_elapsed_ms": elapsed, **json.load(usage)}

cases = [("", 24, "distinct"), ("", 1000, "repeated"), ("", 1000, "distinct"),
         ("-spans", 24, "distinct"), ("-spans", 1000, "distinct"),
         ("-vertical", 24, "distinct"), ("-vertical", 1000, "distinct")]
for pair in range(5):
    for kind, rows, labels in cases:
        order = ["baseline", "candidate"] if pair % 2 == 0 else ["candidate", "baseline"]
        item = {"pair": pair, "kind": kind or "plain", "rows": rows, "labels": labels, "order": order}
        for variant in order:
            item[variant] = run([scratch / (variant + kind), rows, labels, "measured", 5, "compose"])
        record["pairs"].append(item)
    for rows in [24, 1000]:
        for labels in ["repeated", "distinct"]:
            record["headers"].append({"sample": pair, "rows": rows, "labels": labels,
                                       **run([scratch / "headers", rows, labels, "measured", 5, "compose"])})
output.write_text(json.dumps(record, indent=2, sort_keys=True) + "\n")
for kind, rows, labels in cases:
    selected = [x for x in record["pairs"] if (x["kind"], x["rows"], x["labels"]) == (kind or "plain", rows, labels)]
    print(kind or "plain", rows, labels)
    for key in ["plan_cpu_ms", "plan_ms", "compose_cpu_ms", "compose_ms", "process_cpu_ms", "process_elapsed_ms", "rss_kib"]:
        def median(item, variant):
            return item[variant][key] if key in item[variant] else statistics.median(s[key] for s in item[variant]["samples"][1:])
        print(key, "baseline", statistics.median(median(x, "baseline") for x in selected),
              "candidate", statistics.median(median(x, "candidate") for x in selected),
              "paired_percent", statistics.median((median(x, "candidate") / median(x, "baseline") - 1) * 100 for x in selected))
