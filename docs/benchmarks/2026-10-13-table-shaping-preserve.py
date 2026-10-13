"""Compare baseline/candidate geometry, ordered diagnostics, PPTX and SVG.

From repository root: python3 this-file.py SCRATCH. Requires the pinned oracle
environment. Binaries are named {baseline,candidate}-{unmerged,spans}.
Independent reopening also verifies source text, native dimensions and spans.
"""
import json
from pathlib import Path
import subprocess
import sys
import zipfile
from pptx import Presentation

scratch = Path(sys.argv[1])
root = scratch / "preservation"
if root.exists():
    raise SystemExit(f"Use a fresh artifact directory; {root} already exists")
cases = [(24, "distinct", "measured"), (30, "mixed", "measured"),
         (1000, "repeated", "measured"), (1000, "distinct", "measured"),
         (100, "distinct", "measured"), (24, "distinct", "estimated")]
scripts = ["é e\u0301\t42\nExact words", "Δοκιμή κειμένου", "Проверка текста", "مرحبا بالعالم", "日本語の文章"]
record = []
for workload in ["unmerged", "spans"]:
    for count, labels, sizing in cases:
        name = f"{workload}-{count}-{labels}-{sizing}"
        rows = [["ID", "Description", "Value"]]
        for row in range(1, count + 1):
            rows.append(["R" if labels == "repeated" else f"R{row}", scripts[row % len(scripts)] if labels == "mixed" else
                         ("Ready" if count == 100 else "A detailed observation preserves every original word and needs room to wrap naturally."), "1,234"])
        if workload == "spans":
            rows[0] = ["Grouped evidence", "", "Value"]
            for row in range(1, count + 1):
                if row % 3 == 0:
                    rows[row][0] += ": " + rows[row][1]
                    rows[row][1] = ""
        for variant in ["baseline", "candidate"]:
            subprocess.run([str(scratch / f"{variant}-{workload}"), str(count), labels, sizing, "1", "artifacts",
                            str(root / f"{variant}-{name}")], check=True, stdout=subprocess.DEVNULL)
        before, after = [root / f"{v}-{name}" for v in ["baseline", "candidate"]]
        files = sorted(p.name for p in before.iterdir())
        assert files == sorted(p.name for p in after.iterdir())
        for filename in files:
            a, b = (before / filename).read_bytes(), (after / filename).read_bytes()
            # Ignore JSON object ordering only; arrays/diagnostics remain ordered.
            assert (json.loads(a) == json.loads(b)) if filename.endswith(".json") else a == b, (name, filename)
        with zipfile.ZipFile(after / "saved.pptx") as archive:
            assert archive.testzip() is None
        pages = json.loads((after / "geometry.json").read_text())
        assert len(files) == 2 + 2 * (len(pages) + 1)
        assert [r for p in pages for r in p["body_rows"]] == list(range(1, count + 1))
        deck = Presentation(after / "saved.pptx")
        assert len(deck.slides) == len(pages) + 1
        for slide, page in zip(list(deck.slides)[1:], pages):
            native, = [s.table for s in slide.shapes if s.has_table]
            assert [c.width for c in native.columns] == page["widths"]
            assert [r.height for r in native.rows] == page["heights"]
            source_rows = [0] + page["body_rows"]
            assert page["fits"] and len(native.rows) == len(source_rows)
            for local, source in enumerate(source_rows):
                cells = native.rows[local].cells
                assert [c.text.replace("\v", "\n") for c in cells] == rows[source], (name, source)
                merged = workload == "spans" and source % 3 == 0
                assert cells[0].is_merge_origin == merged and cells[1].is_spanned == merged
                assert cells[0].span_width == (2 if merged else 1) and cells[0].span_height == 1
        record.append({"case": name, "files": len(files), "pages": len(pages), "exact": True,
                       "independent_geometry_text_and_topology": True})
print(json.dumps(record, indent=2, sort_keys=True))
