"""Compare two production authored exports and independently reopen the deck.

Usage: python3 this-file.py BASELINE_DIR CANDIDATE_DIR OUTPUT.json
Requires the pinned python-pptx oracle environment.
"""
import base64
import hashlib
import io
import json
from pathlib import Path
import sys
import zipfile
from pptx import Presentation
from PIL import Image

before, after = map(Path, sys.argv[1:3])
files = sorted(p.name for p in before.iterdir())
assert files == sorted(p.name for p in after.iterdir())
for name in files:
    a, b = (before / name).read_bytes(), (after / name).read_bytes()
    assert (json.loads(a) == json.loads(b)) if name.endswith(".json") else a == b, name
with zipfile.ZipFile(after / "saved.pptx") as archive:
    assert archive.testzip() is None
input_deck = json.loads((after / "input.json").read_text())
report = json.loads((after / "report.json").read_text())
deck = Presentation(after / "saved.pptx")
assert len(deck.slides) == report["slide_count"] == 11
assert not report["schema_issues"] and not report["dropped_content"]
groups = {}
for slide in deck.slides:
    name = slide._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}cSld").get("name")
    assert name.startswith("Lectern:")
    source = name.removeprefix("Lectern:").split("-continued-")[0]
    groups.setdefault(source, []).append(slide)
assert list(groups) == [s["id"] for s in input_deck["slides"]]
for source in input_deck["slides"]:
    slides = groups[source["id"]]
    text = "\n".join(shape.text for slide in slides for shape in slide.shapes if shape.has_text_frame)
    body = source["body"]
    words = list(body.get("items", [])) + [b["text"] for b in body.get("bullets", [])]
    for column in ["left", "right"]:
        if column in body:
            words.extend(body[column]["bullets"])
    if body.get("lead"):
        words.append(body["lead"])
    for value in words:
        assert text.encode().count(value.encode()) == 1, (source["id"], value)
    for slide in slides:
        assert source["notes"] in slide.notes_slide.notes_text_frame.text
        if source["id"] == "image":
            # Published native layouts make these picture placeholders; their
            # shape_type is PLACEHOLDER, while the underlying shape stays p:pic.
            picture, = [shape for shape in slide.shapes if shape._element.tag ==
                        "{http://schemas.openxmlformats.org/presentationml/2006/main}pic"]
            blip, = list(picture._element.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"))
            relationship = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
            data = slide.part.related_part(relationship).blob
            assert data == base64.b64decode(
                "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII=")
            Image.open(io.BytesIO(data)).verify()
presentation = deck.part._element
ns = "{http://schemas.microsoft.com/office/powerpoint/2010/main}"
sections = list(presentation.iter(ns + "section"))
assert len(sections) == 1
ids = [int(node.get("id")) for node in sections[0].iter(ns + "sldId")]
assert ids == [slide.slide_id for slide in deck.slides]
record = {"identical_files": len(files), "identical_pptx_and_svgs": True,
          "identical_ordered_diagnostics": True, "slides": len(deck.slides),
          "independent_source_words_notes_images_sections": True,
          "pptx_sha256": hashlib.sha256((after / "saved.pptx").read_bytes()).hexdigest()}
Path(sys.argv[3]).write_text(json.dumps(record, indent=2, sort_keys=True) + "\n")
print(json.dumps(record))
