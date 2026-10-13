"""Paired Release timings or exact saved-output comparisons.

Run from repository root: python3 this-file.py collect|preserve SCRATCH OUTPUT.
SCRATCH contains baseline and candidate binaries compiled from the Swift driver.
The preserve mode needs the pinned python-pptx oracle environment.
"""
import hashlib
import json
import os
from pathlib import Path
import platform
import resource
import statistics
import subprocess
import sys
import tempfile
import zipfile

action, location, output = sys.argv[1:]
scratch = Path(location)

def digest(path):
    result = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(65536), b""):
            result.update(block)
    return result.hexdigest()

source = hashlib.sha256()
for path in sorted(Path("Sources").rglob("*.swift")):
    source.update(str(path).encode() + b"\0")
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(65536), b""):
            source.update(block)
record = {"baseline_commit": "a71153ab96aa9d1877421ed1405b2e88d56b86a0", "candidate_source_sha256": source.hexdigest(),
          "binary_sha256": {v: digest(scratch / v) for v in ["baseline", "candidate"]}}

def run(arguments):
    with tempfile.TemporaryFile() as stdout, tempfile.TemporaryFile() as stderr:
        divisor = 1024 if sys.platform == "darwin" else 1
        parent_peak = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss / divisor
        child = subprocess.Popen([str(a) for a in arguments], stdout=stdout, stderr=stderr)
        _, status, usage = os.wait4(child.pid, 0)
        child.returncode = os.waitstatus_to_exitcode(status)
        stderr.seek(0)
        if child.returncode:
            raise RuntimeError(stderr.read().decode())
        stdout.seek(0)
        return {"samples": json.load(stdout), "peak_rss_kib": usage.ru_maxrss / divisor,
                "collector_peak_rss_kib": parent_peak}

if action == "collect":
    record.update(platform=platform.platform(), swift=subprocess.check_output(["swift", "--version"], text=True))
    cases = [("count", 1000, 8), ("count", 10000, 8), ("count", 10000, 10000),
             ("paragraphs", 4, 180), ("paragraphs", 24, 180), ("paragraphs", 120, 180),
             ("paragraphs", 480, 180), ("text", 30, 180), ("text", 300, 180), ("text", 1200, 180)]
    record["pairs"] = []
    for pair in range(5):
        for mode, count, capacity in cases:
            order = ["baseline", "candidate"] if pair % 2 == 0 else ["candidate", "baseline"]
            sample = {"mode": mode, "count": count, "capacity": capacity, "pair": pair, "order": order}
            for variant in order:
                sample[variant] = run([scratch / variant, mode, count, capacity, 5])
            assert len({trial["pages"] for variant in order for trial in sample[variant]["samples"]}) == 1
            record["pairs"].append(sample)
    for case in cases:
        pairs = [p for p in record["pairs"] if (p["mode"], p["count"], p["capacity"]) == case]
        print(*case)
        for metric in ["cpu_ms", "elapsed_ms", "calls", "visited_items", "peak_rss_kib"]:
            a, b = [], []
            for pair in pairs:
                for values, variant in [(a, "baseline"), (b, "candidate")]:
                    values.append(pair[variant][metric] if metric == "peak_rss_kib" else
                                  statistics.median(x[metric] for x in pair[variant]["samples"] if x["trial"] > 0))
            changes = [(after / before - 1) * 100 for before, after in zip(a, b) if before]
            print(metric, round(statistics.median(a), 4), round(statistics.median(b), 4),
                  "paired %/range", [round(f(changes), 2) for f in [statistics.median, min, max]] if changes else None)
elif action == "preserve":
    from pptx import Presentation
    root = scratch / "preservation"
    if root.exists():
        raise SystemExit(f"Use a fresh artifact directory: {root}")
    record["cases"] = []
    for mode, count, height in [("paragraphs", 4, 180), ("paragraphs", 24, 180),
                                ("paragraphs", 120, 180), ("paragraphs", 120, 260),
                                ("text", 30, 180), ("text", 300, 180)]:
        name = f"{mode}-{count}-{height}"
        for variant in ["baseline", "candidate"]:
            subprocess.run([str(scratch / variant), mode, str(count), str(height), "1",
                            str(root / f"{variant}-{name}")], check=True, stdout=subprocess.DEVNULL)
        before, after = [root / f"{variant}-{name}" for variant in ["baseline", "candidate"]]
        files = sorted(p.name for p in before.iterdir())
        assert files == sorted(p.name for p in after.iterdir())
        for filename in files:
            a, b = (before / filename).read_bytes(), (after / filename).read_bytes()
            assert (json.loads(a) == json.loads(b)) if filename.endswith(".json") else a == b, (name, filename)
        pages = json.loads((after / "pages.json").read_text())
        with zipfile.ZipFile(after / "saved.pptx") as archive:
            assert archive.testzip() is None
        deck = Presentation(after / "saved.pptx")
        assert len(deck.slides) == len(pages) + 1
        for index, (slide, words) in enumerate(zip(list(deck.slides)[1:], pages)):
            shapes = [s for s in slide.shapes if s.has_text_frame]
            title, body = shapes
            assert title.text == f"Evidence — page {index + 1}"
            assert body.text.replace("\v", "\n").encode() == "\n".join(words).encode()
            assert (body.left, body.top, body.width, body.height) == tuple(p * 12700 for p in [48, 96, 520, height])
            assert "Original words and ordering retained." in slide.notes_slide.notes_text_frame.text
        record["cases"].append({"case": name, "pages": len(pages), "files": len(files),
                                "identical": True, "independent_text_geometry_notes": True})
else:
    raise SystemExit("Expected collect or preserve")
Path(output).write_text(json.dumps(record, indent=2, sort_keys=True) + "\n")
