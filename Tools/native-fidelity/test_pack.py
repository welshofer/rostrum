#!/usr/bin/env python3
"""Integration checks against a built executable; no native acceptance implied."""
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import xml.etree.ElementTree as ET
import zipfile
from pptx import Presentation

from run import collect, command_output

ROOT = Path(__file__).resolve().parents[2]
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}


class FidelityPackTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.binary = Path(os.environ.get("ROSTRUM_NATIVE_FIDELITY_BINARY", ROOT / ".build/debug/native-fidelity"))
        if not cls.binary.is_file():
            raise RuntimeError("Build native-fidelity before running this suite")
        cls.temp = tempfile.TemporaryDirectory()
        cls.addClassCleanup(cls.temp.cleanup)
        cls.directory = Path(cls.temp.name)
        fixture = ROOT / "Tests/RostrumTests/Fixtures/NativeGlyphPlacement/eligibility/fonts"
        cls.config = {"family": "DejaVu Sans", "regular": str(fixture / "DejaVuSans.ttf"),
                      "bold": str(fixture / "DejaVuSans-Bold.ttf"), "italic": str(fixture / "DejaVuSans-Oblique.ttf")}
        cls.fonts = cls.directory / "fonts.json"
        cls.fonts.write_text(json.dumps(cls.config))
        cls.output = cls.directory / "pack"
        cls.lectern = Path(os.environ["ROSTRUM_LECTERN_INSPECTION_BINARY"]) if os.environ.get("ROSTRUM_LECTERN_INSPECTION_BINARY") else None
        cls.manifest = collect(cls.binary, cls.fonts, cls.output, runs=2, build_mode="unspecified", lectern_binary=cls.lectern)

    def test_four_profiles_repeat_identically_but_stay_unreviewed(self):
        self.assertTrue(self.manifest["complete"])
        self.assertEqual(len(self.manifest["profiles"]), 4)
        for key in ("powerpointSlideView", "powerpointExport"):
            self.assertEqual(self.manifest[key], "pending")
        self.assertEqual(self.manifest["lecternInspect"],
                         "core-artifacts-collected; native-review-pending" if self.lectern else "pending")
        for profile in self.manifest["profiles"]:
            first, second = profile["samples"]
            self.assertEqual(first["artifactSHA256"], second["artifactSHA256"])
            self.assertGreater(first["peakRSSBytes"], 0)
            self.assertGreater(first["userCPUSeconds"] + first["systemCPUSeconds"], 0)

    def test_native_content_geometry_and_font_storage(self):
        for profile in self.manifest["profiles"]:
            case = self.output / profile["samples"][0]["directory"]
            report = json.loads((case / "report.json").read_text())
            with zipfile.ZipFile(case / "deck.pptx") as deck, zipfile.ZipFile(case / "input.potx") as template:
                size = ET.fromstring(deck.read("ppt/presentation.xml")).find("p:sldSz", NS)
                self.assertEqual(int(size.attrib["cy"]), 540 * 12700)
                self.assertEqual(int(size.attrib["cx"]), (720 if report["aspect"] == "4x3" else 960) * 12700)
                embedded = ET.fromstring(deck.read("ppt/presentation.xml")).find("p:embeddedFontLst", NS)
                self.assertEqual(embedded is not None, report["fontMode"] == "embedded")
                if embedded is not None:
                    for role in ("regular", "bold", "italic"):
                        self.assertIsNotNone(embedded.find(f"p:embeddedFont/p:{role}", NS))
                for name in template.namelist():
                    if name.startswith(("ppt/slideLayouts/", "ppt/slideMasters/", "ppt/theme/")):
                        self.assertEqual(template.read(name), deck.read(name), name)
                first_slide = ET.fromstring(deck.read("ppt/slides/slide1.xml"))
                self.assertIn("Rostrum in Lectern", "".join(first_slide.itertext()))
                self.assertEqual(len(first_slide.findall(".//a:br", NS)), 2)
                self.assertIn("Indented line retained", "".join(first_slide.itertext()))
                styles = [node.attrib for node in first_slide.findall(".//a:rPr", NS)]
                self.assertTrue(any(p.get("b") == "1" for p in styles))
                self.assertTrue(any(p.get("i") == "1" for p in styles))
                self.assertTrue(any(p.get("sz") == "2200" for p in styles))
                self.assertEqual([row for page in report["pages"] for row in page["sourceRows"]], list(range(1, 25)))
                for page in report["pages"]:
                    root = ET.fromstring(deck.read(f"ppt/slides/slide{page['slide']}.xml"))
                    table = root.find(".//a:tbl", NS)
                    rows = table.findall("a:tr", NS)
                    self.assertEqual(rows[0].find("a:tc", NS).attrib.get("gridSpan"), "2")
                    for native, source in zip(rows[1:], page["sourceRows"]):
                        texts = ["".join(cell.itertext()) for cell in native.findall("a:tc", NS)]
                        self.assertEqual(texts, [f"Record {source}",
                            f"Exact source evidence {source} stays editable and wraps at full size.", "Reviewed"])
                    self.assertEqual([int(row.attrib["h"]) for row in rows], page["rowHeightsEMU"])

    def test_refuses_existing_output(self):
        with self.assertRaisesRegex(ValueError, "already exists"):
            collect(self.binary, self.fonts, self.output)
        result = subprocess.run([str(self.binary), str(self.fonts), str(self.output), "4x3", "external"], capture_output=True, text=True)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("already exists", result.stderr)

    def test_merged_groups_keep_native_topology_content_and_geometry(self):
        for profile in self.manifest["profiles"]:
            case = self.output / profile["samples"][0]["directory"]
            report = json.loads((case / "report.json").read_text())
            source = json.loads((case / "merged-source.json").read_text())
            pages = json.loads((case / "merged-geometry.json").read_text())
            self.assertEqual(pages, report["mergedPages"])
            self.assertGreater(len(pages), 1)
            self.assertEqual([row for page in pages for row in page["sourceRows"]], list(range(1, 25)))
            deck = Presentation(case / "deck.pptx")
            for page in pages:
                first, last = page["sourceRows"][0], page["sourceRows"][-1]
                self.assertEqual((first - 1) % 4, 0)
                self.assertEqual(last % 4, 0)
                table, = [shape.table for shape in deck.slides[page["slide"] - 1].shapes if shape.has_table]
                self.assertEqual([row.height for row in table.rows], page["rowHeightsEMU"])
                self.assertEqual([column.width for column in table.columns], page["columnWidthsEMU"])
                self.assertEqual([[cell.text for cell in row.cells] for row in table.rows],
                                 [source["rows"][0]] + [source["rows"][i] for i in page["sourceRows"]])
                expected = {(m["row"] - first + 1, m["column"]): (m["rowSpan"], m["columnSpan"])
                            for m in source["merges"] if first <= m["row"] <= last}
                actual = {(r, c): (table.cell(r, c).span_height, table.cell(r, c).span_width)
                          for r in range(len(table.rows)) for c in range(len(table.columns))
                          if table.cell(r, c).is_merge_origin}
                self.assertEqual(actual, expected)

    def assert_invalid_font(self, change, message):
        with tempfile.TemporaryDirectory(dir=self.directory) as temporary:
            directory = Path(temporary)
            config = dict(self.config, **change)
            path = directory / "fonts.json"
            path.write_text(json.dumps(config))
            output = directory / "result"
            result = subprocess.run([str(self.binary), str(path), str(output), "4x3", "external"], capture_output=True, text=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn(message, result.stderr)
            self.assertFalse(output.exists())

    def test_face_substitution_is_rejected(self):
        self.assert_invalid_font({"regular": self.config["bold"]}, "expected exact family")

    def test_family_substitution_is_rejected(self):
        self.assert_invalid_font({"family": "Calibri"}, "expected exact family")

    def test_missing_face_is_rejected(self):
        self.assert_invalid_font({"italic": str(self.directory / "absent.ttf")}, "native-fidelity:")

    @unittest.skipUnless(os.environ.get("ROSTRUM_LECTERN_INSPECTION_BINARY"), "Lectern executable not supplied")
    def test_lectern_inspects_and_exports_every_source_row_note_and_section(self):
        for profile in self.manifest["profiles"]:
            first, second = [sample["lectern"] for sample in profile["samples"]]
            self.assertEqual(first["artifactSHA256"], second["artifactSHA256"])
            self.assertGreater(first["peakRSSBytes"], 0)
            case = self.output / first["directory"]
            evidence = json.loads((case / "inspection.json").read_text())
            markdown = (case / "export/deck/deck.md").read_text()
            total = evidence["slideCount"]
            self.assertEqual(evidence["export"]["slideCount"], total)
            self.assertEqual([p["number"] for p in evidence["previews"]], list(range(1, total + 1)))
            self.assertTrue(all(p["rendered"] for p in evidence["previews"]))
            for preview in evidence["previews"]:
                self.assertAlmostEqual(preview["width"] / preview["height"],
                                       4 / 3 if profile["name"].startswith("4x3") else 16 / 9)
            pack_case = self.output / profile["samples"][0]["directory"]
            pack_report = json.loads((pack_case / "report.json").read_text())
            merged_start = pack_report["mergedPages"][0]["slide"] - 1
            self.assertEqual([s["name"] for s in evidence["sections"]], ["Typography", "Tables", "Merged tables"])
            self.assertEqual([s["slideIndices"] for s in evidence["sections"]],
                             [[0], list(range(1, merged_start)), list(range(merged_start, total))])
            rows = [row for slide in evidence["slides"][:merged_start] for table in slide["tables"] for row in table["rows"][1:]]
            expected = [[f"Record {i}", f"Exact source evidence {i} stays editable and wraps at full size.", "Reviewed"]
                        for i in range(1, 25)]
            self.assertEqual(rows, expected)
            merged_rows = [row for slide in evidence["slides"][merged_start:] for table in slide["tables"] for row in table["rows"][1:]]
            merged_source = json.loads((pack_case / "merged-source.json").read_text())
            self.assertEqual(merged_rows, merged_source["rows"][1:])
            for row in merged_rows:
                for text in row:
                    if text:
                        self.assertIn(text, markdown)
            for row in expected:
                self.assertIn(" | ".join(row), markdown)
            for slide in evidence["slides"]:
                note = f"Native fidelity fixture: slide {slide['number']}. Preserve these exact notes."
                self.assertIn(note, slide["notes"])
                self.assertIn(note, markdown)
            self.assertIn("Rostrum in Lectern", markdown)
            self.assertIn("Indented line retained", markdown)
            self.assertEqual(evidence["slides"][0]["bullets"], [
                "Inherited font, tabs and soft breaks",
                "Alpha\t12.5 Beta\t38.0 Gamma\t7.25",
                "    Indented line retained",
                "regular: AV office café. bold: AV office café. italic: AV office café."])
            for text in ("Alpha\t12.5", "Beta\t38.0", "Gamma\t7.25",
                         "regular: AV office café. bold: AV office café. italic: AV office café."):
                self.assertIn(text, markdown)
            self.assertIn("## Sections", markdown)
            self.assertIn("Typography", markdown)
            self.assertIn("Tables", markdown)
            self.assertEqual(evidence["schemaIssues"], [])
            self.assertEqual(evidence["readWarnings"], [])
            self.assertEqual(evidence["outlineWarnings"], [])
            codes = [issue["code"] for diagnostic in evidence["previewDiagnostics"] for issue in diagnostic["issues"]]
            if "-embedded" in profile["name"]:
                self.assertNotIn("missingFont", codes)
            elif sys.platform == "linux":
                self.assertIn("missingFont", codes)
            timing = first["timings"]
            self.assertEqual(timing["nativeVisualAcceptance"], "pending")
            self.assertTrue(timing["firstPreviewRendered"])
            self.assertEqual(timing["stages"]["droppedSamples"], 0)
            stages = timing["stages"]["samples"]
            self.assertEqual([s["slideNumber"] for s in stages if s["stage"] == "preview.svg"], list(range(1, total + 1)))
            self.assertTrue({"inspect.read", "inspect.open", "inspect.extract", "preview.embeddedFonts",
                             "preview.fontDiscovery", "preview.installedFonts"}.issubset({s["stage"] for s in stages}))
            self.assertGreater(timing["firstPreviewSlotCompleted"]["elapsedSeconds"], 0)
            self.assertLessEqual(timing["firstPreviewSlotCompleted"]["elapsedSeconds"], timing["inspection"]["elapsedSeconds"])

    @unittest.skipUnless(os.environ.get("ROSTRUM_LECTERN_INSPECTION_BINARY"), "Lectern executable not supplied")
    def test_lectern_refuses_existing_output_and_missing_input(self):
        deck = self.output / self.manifest["profiles"][0]["samples"][0]["directory"] / "deck.pptx"
        result = subprocess.run([str(self.lectern), str(deck), str(self.output)], capture_output=True, text=True)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("already exists", result.stderr)
        absent_output = self.directory / "absent-result"
        result = subprocess.run([str(self.lectern), str(self.directory / "absent.pptx"), str(absent_output)],
                                capture_output=True, text=True)
        self.assertNotEqual(result.returncode, 0)
        self.assertFalse(absent_output.exists())


class ProvenanceTests(unittest.TestCase):
    def test_container_checkout_ownership(self):
        # Git's own hook simulates the host/container owner mismatch seen in CI.
        with patch.dict(os.environ, {"GIT_TEST_ASSUME_DIFFERENT_OWNER": "1"}):
            self.assertRegex(command_output("git", "rev-parse", "HEAD"), r"^[0-9a-f]{40}$")
            command_output("git", "status", "--porcelain")
            with self.assertRaisesRegex(RuntimeError, "fatal:"):
                command_output("git", "rev-parse", "--verify", "refs/heads/absent-fidelity-provenance-probe")


if __name__ == "__main__":
    program = unittest.main(exit=False)
    if os.environ.get("GITHUB_ACTIONS") == "true":
        # Class setup failures can run zero tests. Publish their tracebacks as
        # check annotations too, so an exit code alone cannot hide the cause.
        for test, detail in program.result.errors + program.result.failures:
            message = f"{test}\n{detail}".replace("%", "%25").replace("\r", "%0D").replace("\n", "%0A")
            print(f"::error title=Native fidelity pack::{message}", flush=True)
    sys.exit(0 if program.result.wasSuccessful() and program.result.testsRun else 1)
