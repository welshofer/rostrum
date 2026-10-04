#!/usr/bin/env python3
"""Create an owned two-image slide with 4,096 relationships for sparse lookup QA.

Requires python-pptx 1.0.2. The unused arcs are external hyperlinks, never fetched.
The two distinct image arcs occur near the start; indexing all arcs would waste
work. Refuses to replace an existing fixture. Original test content is CC0.
"""
import argparse
import base64
from io import BytesIO
from pathlib import Path
import zipfile
from xml.etree import ElementTree as ET

import pptx
from pptx import Presentation
from pptx.util import Inches

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('output', type=Path)
args = parser.parse_args()
if pptx.__version__ != '1.0.2':
    parser.error('python-pptx 1.0.2 required')
if args.output.exists():
    parser.error('output exists')
png = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII=')
deck = Presentation()
deck.slide_width, deck.slide_height = Inches(12), Inches(7)
slide = deck.slides.add_slide(deck.slide_layouts[6])
for number in range(2):
    # Valid PNG with distinct trailing bytes prevents python-pptx media/rId reuse.
    slide.shapes.add_picture(BytesIO(png + bytes([number])), Inches(1 + number * 3), Inches(1), Inches(2), Inches(2))
saved = BytesIO()
deck.save(saved)
ns = 'http://schemas.openxmlformats.org/package/2006/relationships'
ET.register_namespace('', ns)
with zipfile.ZipFile(saved) as source:
    entries = {name: source.read(name) for name in source.namelist()}
name = 'ppt/slides/_rels/slide1.xml.rels'
rels = ET.fromstring(entries[name])
for number in range(len(rels) + 1, 4097):
    ET.SubElement(rels, '{' + ns + '}Relationship', {
        'Id': f'rId{number}',
        'Type': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
        'Target': f'https://example.invalid/unused/{number}', 'TargetMode': 'External'})
assert len(rels) == 4096
entries[name] = ET.tostring(rels, encoding='utf-8', xml_declaration=True)
args.output.parent.mkdir(parents=True, exist_ok=True)
with zipfile.ZipFile(args.output, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
    for name, data in sorted(entries.items()):
        entry = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
        entry.compress_type = zipfile.ZIP_DEFLATED
        archive.writestr(entry, data)
print(args.output)
