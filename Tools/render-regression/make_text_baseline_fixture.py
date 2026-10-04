#!/usr/bin/env python3
"""Author the CC0 three-slide typography fixture with python-pptx 1.0.2.

Run once to create fixtures/text-baseline-v1.pptx and its coordinate manifest.
Run with --verify to regenerate in memory, reopen, and compare exact bytes.
No fonts or Office-rendered references are generated or embedded. Existing
artifacts are never overwritten; use a new --id for a changed fixture.
"""
import argparse
from datetime import datetime
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
import zipfile

import pptx
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

EMU_PER_INCH = 914400
MARGINS = {'marL': 100584, 'marT': 64008, 'marR': 210312, 'marB': 118872}
NO_GRID = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'
FONTS = [
    {'family': 'Arial', 'bold': False, 'fileName': 'Arial.ttf',
     'sha256': '525979822591a3447cfc49d943d6f7683508e25543407871c0ed8fed05fd2bd9'},
    {'family': 'Arial', 'bold': True, 'fileName': 'Arial Bold.ttf',
     'sha256': 'd72db21f9242aedd6b917d8549ad5921766b24d5f8d0becfda2ff4c620b3c2e0'},
    {'family': 'Calibri', 'bold': False, 'fileName': 'Calibri.ttf',
     'sha256': 'ea801e1f869b55464339058b1d4263d07cc074a18e20aa3ee1d07901423dee53'},
    {'family': 'Calibri', 'bold': True, 'fileName': 'Calibrib.ttf',
     'sha256': 'ac1cf97565de97cdc322228d875dc18c1131656c5138173e2c6d8ac7a37aa7f2'}]


def require(condition, message):
    if not condition:
        raise ValueError(message)


def run(text, size, bold=False):
    return {'kind': 'run', 'text': text, 'pointSize': size, 'bold': bold}


def line_break(size):
    return {'kind': 'break', 'pointSize': size, 'bold': False}


def case(identifier, family, frame, runs, wrap=False, spacing='absent', default_size=None):
    return {'id': identifier, 'family': family,
            'frameEMU': [int(Inches(value)) for value in frame],
            'wrap': 'square' if wrap else 'none', 'lineSpacing': spacing,
            'paragraphDefaultPointSize': runs[0]['pointSize'] if default_size is None else default_size,
            'paragraphDefaultBold': runs[0]['bold'],
            'runs': runs,
            'authoredText': ''.join(r.get('text', '\v') for r in runs)}


def cases():
    slides = [[], [], []]
    faces = [('Arial', False), ('Arial', True), ('Calibri', False), ('Calibri', True)]
    for row, size in enumerate([14, 18, 24]):
        for column, (family, bold) in enumerate(faces):
            identifier = f's1-{family.lower()}-{"bold" if bold else "regular"}-{size}'
            slides[0].append(case(identifier, family, (.5 + column * 2.75, .5 + row * 2, 2.5, 1.5),
                                  [run('Hpxgy', size, bold)]))
    variants = [('Arial', 'absent'), ('Arial', 'explicit100percent'),
                ('Calibri', 'absent'), ('Calibri', 'explicit100percent')]
    for row, (first, second) in enumerate([(24, 14), (14, 24), (18, 18)]):
        for column, (family, spacing) in enumerate(variants):
            identifier = f's2-{family.lower()}-{first}-{second}-{spacing}'
            slides[1].append(case(identifier, family, (.5 + column * 2.75, .5 + row * 2, 2.5, 1.5),
                                  [run('Hpxgy', first), line_break(second), run('Hpxgy', second)], spacing=spacing))
    for row, family in enumerate(['Arial', 'Calibri']):
        for column, mode in enumerate(['mixed-auto', 'mixed-explicit-break', 'all18-auto']):
            runs = [run('Mixed ', 18), run('bold', 18 if mode == 'all18-auto' else 24, True)]
            if mode == 'mixed-explicit-break':
                runs += [line_break(14), run('text', 14)]
            else:
                runs += [run(' text', 18 if mode == 'all18-auto' else 14)]
            slides[2].append(case(f's3-{family.lower()}-{mode}', family, (1 + column * 3.5, 1 + row * 2.75, 2, 2),
                                  runs, wrap=True))
    return slides


def properties(tag, family, size, bold):
    pr = OxmlElement(tag)
    for name, value in {'sz': str(size * 100), 'b': '1' if bold else '0', 'i': '0', 'u': 'none',
                        'spc': '0', 'kern': '1200', 'baseline': '0', 'lang': 'en-US'}.items():
        pr.set(name, value)
    fill = OxmlElement('a:solidFill')
    color = OxmlElement('a:srgbClr'); color.set('val', '000000'); fill.append(color); pr.append(fill)
    for tag in ['a:latin', 'a:ea', 'a:cs']:
        face = OxmlElement(tag); face.set('typeface', family); pr.append(face)
    return pr


def add_case(slide, record):
    shape = slide.shapes.add_table(1, 1, *record['frameEMU'])
    shape.name = record['id']
    table = shape.table
    table.rows[0].height = record['frameEMU'][3]
    table.columns[0].width = record['frameEMU'][2]
    flags = table._tbl.tblPr
    for name in ['firstRow', 'lastRow', 'firstCol', 'lastCol', 'bandRow', 'bandCol', 'rtl']:
        flags.set(name, '0')
    flags.find(qn('a:tableStyleId')).text = NO_GRID
    cell = table.cell(0, 0)
    cell.margin_left, cell.margin_top = MARGINS['marL'], MARGINS['marT']
    cell.margin_right, cell.margin_bottom = MARGINS['marR'], MARGINS['marB']
    cell.vertical_anchor = MSO_ANCHOR.TOP
    cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(255, 255, 255)
    tcpr = cell._tc.get_or_add_tcPr()
    tcpr.set('vert', 'horz'); tcpr.set('anchorCtr', '0')
    for index, edge in enumerate(['lnL', 'lnR', 'lnT', 'lnB', 'lnTlToBr', 'lnBlToTr']):
        line = OxmlElement('a:' + edge); line.set('w', '0'); line.append(OxmlElement('a:noFill'))
        tcpr.insert(index, line)
    frame = cell.text_frame
    frame.clear(); frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.word_wrap = record['wrap'] != 'none'
    body = frame._txBody.bodyPr
    body.set('anchor', 't'); body.set('vert', 'horz'); body.set('rtlCol', '0')
    # Table margins are authoritative; body insets repeat them to eliminate an
    # unrelated inheritance variable when an Office consumer inspects the text.
    for margin, inset in [('marL', 'lIns'), ('marT', 'tIns'), ('marR', 'rIns'), ('marB', 'bIns')]:
        body.set(inset, str(MARGINS[margin]))
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    paragraph.space_before = Pt(0); paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1.0 if record['lineSpacing'] == 'explicit100percent' else None
    ppr = paragraph._p.get_or_add_pPr()
    for name, value in {'marL': '0', 'marR': '0', 'indent': '0', 'lvl': '0', 'rtl': '0'}.items():
        ppr.set(name, value)
    ppr.append(OxmlElement('a:buNone'))
    ppr.append(properties('a:defRPr', record['family'], record['paragraphDefaultPointSize'], record['paragraphDefaultBold']))
    for token in record['runs']:
        element = OxmlElement('a:r' if token['kind'] == 'run' else 'a:br')
        element.append(properties('a:rPr', record['family'], token['pointSize'], token['bold']))
        if token['kind'] == 'run':
            text = OxmlElement('a:t'); text.text = token['text']; element.append(text)
        paragraph._p.append(element)
    record['shapeId'] = shape.shape_id
    record['textOriginXEMU'] = record['frameEMU'][0] + MARGINS['marL']
    record['textContentTopEMU'] = record['frameEMU'][1] + MARGINS['marT']
    record['availableTextWidthEMU'] = record['frameEMU'][2] - MARGINS['marL'] - MARGINS['marR']


def canonical_package(raw):
    target = BytesIO()
    with zipfile.ZipFile(BytesIO(raw)) as source, zipfile.ZipFile(target, 'w') as output:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, (1980, 1, 1, 0, 0, 0))
            info.create_system = 3; info.external_attr = 0o600 << 16
            info.compress_type = zipfile.ZIP_STORED
            output.writestr(info, source.read(name))
    return target.getvalue()


def verify(payload, records):
    with zipfile.ZipFile(BytesIO(payload)) as package:
        require(package.testzip() is None, 'ZIP CRC failure')
        require(not any(name.startswith('ppt/fonts/') for name in package.namelist()), 'Unexpected embedded font')
    deck = Presentation(BytesIO(payload))
    require((deck.slide_width, deck.slide_height) == (Inches(12), Inches(7)), 'Slide size differs')
    require(len(deck.slides) == 3, 'Expected exactly three slides')
    shape_count = run_count = break_count = 0
    for slide, expected in zip(deck.slides, records):
        require(len(slide.shapes) == len(expected), 'Shape count differs')
        for shape, record in zip(slide.shapes, expected):
            require(shape.has_table and shape.name == record['id'] and shape.shape_id == record['shapeId'], 'Shape identity differs')
            require([shape.left, shape.top, shape.width, shape.height] == record['frameEMU'], 'Authored frame differs')
            table = shape.table
            require(len(table.rows) == len(table.columns) == 1, 'Expected single-cell table')
            require(table.rows[0].height == shape.height and table.columns[0].width == shape.width, 'Grid dimensions differ')
            require(table._tbl.tblPr.find(qn('a:tableStyleId')).text == NO_GRID, 'Table style differs')
            cell = table.cell(0, 0); tcpr = cell._tc.tcPr; frame = cell.text_frame
            require(all(tcpr.get(key) == str(value) for key, value in MARGINS.items()), 'Cell margins differ')
            require(tcpr.get('anchor') == 't' and tcpr.get('vert') == 'horz', 'Cell anchoring differs')
            require(all(tcpr.find(qn('a:' + edge)).find(qn('a:noFill')) is not None
                        for edge in ['lnL', 'lnR', 'lnT', 'lnB', 'lnTlToBr', 'lnBlToTr']), 'Cell border differs')
            require(str(cell.fill.fore_color.rgb) == 'FFFFFF', 'Cell fill differs')
            require(frame.auto_size == MSO_AUTO_SIZE.NONE and frame._txBody.bodyPr.get('wrap') == record['wrap'], 'Autofit/wrap differs')
            require(len(frame.paragraphs) == 1 and frame.text == record['authoredText'], 'Authored text differs')
            paragraph = frame.paragraphs[0]; ppr = paragraph._p.pPr
            require(paragraph.alignment == PP_ALIGN.LEFT and paragraph.space_before == paragraph.space_after == 0, 'Paragraph spacing/alignment differs')
            require(all(ppr.get(name) == '0' for name in ['marL', 'marR', 'indent', 'lvl', 'rtl']), 'Paragraph geometry differs')
            default = ppr.find(qn('a:defRPr'))
            require(default.get('sz') == str(record['paragraphDefaultPointSize'] * 100)
                    and default.get('b') == str(int(record['paragraphDefaultBold'])), 'Paragraph default size/weight differs')
            require(all(default.find(qn('a:' + script)).get('typeface') == record['family']
                        for script in ['latin', 'ea', 'cs']), 'Paragraph default face differs')
            spacing = ppr.find(qn('a:lnSpc'))
            require((spacing is None) if record['lineSpacing'] == 'absent'
                    else spacing is not None and spacing.find(qn('a:spcPct')).get('val') == '100000', 'Line spacing differs')
            actual = [node for node in paragraph._p if node.tag in [qn('a:r'), qn('a:br')]]
            require(len(actual) == len(record['runs']), 'Run/break count differs')
            for element, token in zip(actual, record['runs']):
                require(element.tag == qn('a:r' if token['kind'] == 'run' else 'a:br'), 'Run/break order differs')
                rpr = element.find(qn('a:rPr'))
                require(rpr.get('sz') == str(token['pointSize'] * 100) and rpr.get('b') == str(int(token['bold'])), 'Font size/weight differs')
                require(all(rpr.find(qn('a:' + script)).get('typeface') == record['family'] for script in ['latin', 'ea', 'cs']), 'Font face differs')
                require(rpr.find(qn('a:solidFill')).find(qn('a:srgbClr')).get('val') == '000000', 'Text color differs')
                if token['kind'] == 'run':
                    require(element.find(qn('a:t')).text == token['text'], 'Run text differs'); run_count += 1
                else:
                    break_count += 1
            shape_count += 1
    return {'slides': 3, 'shapesPerSlide': [12, 12, 6], 'singleCellTables': shape_count,
            'textRuns': run_count, 'explicitBreaks': break_count, 'reopenedWith': 'python-pptx 1.0.2',
            'zipCRCValid': True, 'embeddedFonts': False, 'officeAcceptance': 'not performed'}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--id', default='text-baseline-v1')
    parser.add_argument('--output-dir', type=Path, default=Path(__file__).with_name('fixtures'))
    parser.add_argument('--verify', action='store_true')
    args = parser.parse_args()
    require(pptx.__version__ == '1.0.2', 'python-pptx 1.0.2 is required')
    require(re.fullmatch('[a-z0-9-]+', args.id), 'Use lowercase letters, digits and hyphens for --id')
    deck = Presentation(); deck.slide_width = Inches(12); deck.slide_height = Inches(7)
    core = deck.core_properties
    core.title = args.id; core.subject = 'Direct-export typography measurements'
    core.author = core.last_modified_by = 'Rostrum conformance'
    core.comments = 'Original project-authored fixture content, CC0. No font binaries embedded.'
    core.created = core.modified = datetime(2026, 10, 2); core.revision = 1
    records = cases()
    for cases_on_slide in records:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        for record in cases_on_slide:
            add_case(slide, record)
    stream = BytesIO(); deck.save(stream)
    payload = canonical_package(stream.getvalue())
    verification = verify(payload, records)
    manifest = {'schema': 1, 'id': args.id, 'file': args.id + '.pptx',
                'sha256': hashlib.sha256(payload).hexdigest(), 'producer': 'python-pptx 1.0.2',
                'license': 'CC0-1.0 for original project-authored fixture content and generator',
                'provenance': 'Independent producer fixture; standard python-pptx template packaging retained. No Rostrum serialization, font binaries, or reference images used.',
                'generatorSHA256': hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
                'reproducibility': 'Fixed core timestamps; sorted ZIP_STORED entries with fixed ZIP timestamps/attributes.',
                'slideDimensionsEMU': [int(Inches(12)), int(Inches(7))], 'referenceFonts': FONTS,
                'common': {'cellMarginsEMU': MARGINS, 'bodyInsetsRepeatCellMargins': True, 'verticalAnchor': 'top',
                           'cellStyle': NO_GRID, 'borders': 'explicit noFill on all six edges', 'cellFill': 'FFFFFF',
                           'textColor': '000000', 'autofit': 'noAutofit', 'paragraphSpaceBeforeAfterPoints': 0,
                           'paragraphDefaultPointSize': 'per case: single-run size, transition first size, mixed 18pt',
                           'paragraphDefaultBold': 'per case: first run weight', 'trackingPoints': 0,
                           'kerningThresholdPoints': 12, 'italic': False, 'underline': 'none', 'language': 'en-US',
                           'endParagraphRunProperties': 'absent'},
                'slides': [{'slideIndex': i + 1, 'name': name, 'cases': content} for i, (name, content) in enumerate(zip(
                    ['single-line-metrics', 'line-transition-metrics', 'mixed-run-wrap-controls'], records))],
                'verification': verification,
                'referenceStatus': 'awaiting independent Office exports; no visual equivalence established',
                'independentReferences': {'directSlidePDF': None, 'png1200x700': [None] * 3, 'png2400x1400': [None] * 3},
                'captureInstructions': ['Export all three slides directly to a PDF preserving 12x7-inch slide dimensions; record Office version and route.',
                                        'Export each slide as PNG at 1200x700 and 2400x1400; retain original fixture and reference bytes with hashes.',
                                        'Verify exported PDF glyph outlines against the pinned local fonts. No font bytes belong in the repository.',
                                        'Do not infer baselines or goldens from this authored-coordinate manifest; capture through the coordinated Office-session owner.']}
    encoded = json.dumps(manifest, indent=2) + '\n'
    path = args.output_dir / (args.id + '.pptx'); metadata = args.output_dir / (args.id + '.manifest.json')
    if args.verify:
        require(path.read_bytes() == payload, 'Fixture bytes differ from deterministic regeneration')
        require(metadata.read_text() == encoded, 'Manifest differs from deterministic regeneration')
    else:
        require(not path.exists() and not metadata.exists(), 'Fixture exists; use --verify or choose a new --id')
        args.output_dir.mkdir(parents=True, exist_ok=True)
        path.write_bytes(payload); metadata.write_text(encoded)
    print(json.dumps({'mode': 'verified' if args.verify else 'created', 'file': str(path),
                      'sha256': manifest['sha256'], 'verification': verification}, indent=2))


if __name__ == '__main__':
    main()
