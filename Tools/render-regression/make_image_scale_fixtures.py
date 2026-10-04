#!/usr/bin/env python3
"""Make deterministic subsets of the pinned dense image fixture, without overwrites.

python3 Tools/render-regression/make_image_scale_fixtures.py --output-dir /tmp/image-scale

The 2,000-picture case remains the original fixture. Subsets retain its first
N pictures, geometry, image bytes and relationship order. Only trailing picture
elements, their relationships/media and any matching content-type overrides
are removed. This is synthetic scaling evidence, not visual conformance.
"""
import argparse
import hashlib
from io import BytesIO
import json
from pathlib import Path
import platform
import posixpath
import zipfile
import zlib

from lxml import etree
import pptx
from pptx import Presentation

SOURCE_SHA256 = '4b56a1c7ac25daf389cb37c695d865e22746d9e6c6559f70b5b845fed42a880c'
COUNTS = (100, 250, 510, 511, 768, 1024)
SLIDE = 'ppt/slides/slide1.xml'
RELS = 'ppt/slides/_rels/slide1.xml.rels'
P = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
REL_TYPE = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/'


def digest(data):
    return hashlib.sha256(data).hexdigest()


def parse(data):
    return etree.fromstring(data, parser=etree.XMLParser(resolve_entities=False, no_network=True))


def serialize(element):
    return etree.tostring(element, encoding='UTF-8', xml_declaration=True, standalone=True)


def media_target(relationship):
    assert relationship.get('TargetMode') != 'External'
    target = posixpath.normpath(posixpath.join('ppt/slides', relationship.get('Target')))
    assert target.startswith('ppt/media/')
    return target


def verify(data, count):
    presentation = Presentation(BytesIO(data))
    assert len(presentation.slides) == 1
    assert len(presentation.slides[0].shapes) == count
    with zipfile.ZipFile(BytesIO(data)) as archive:
        pictures = parse(archive.read(SLIDE)).find(P + 'cSld').find(P + 'spTree').findall(P + 'pic')
        relationships = list(parse(archive.read(RELS)))
        images = [rel for rel in relationships if rel.get('Type') == REL_TYPE + 'image']
        layout = [rel for rel in relationships if rel.get('Type') == REL_TYPE + 'slideLayout']
        assert len(pictures) == len(images) == count
        assert len(layout) == 1 and len(relationships) == count + 1
        ids = {rel.get('Id') for rel in images}
        assert len(ids) == count
        assert {pic.find('.//' + A + 'blip').get(R + 'embed') for pic in pictures} == ids
        targets = {media_target(rel) for rel in images}
        media = {name for name in archive.namelist() if name.startswith('ppt/media/')}
        assert targets == media and len(media) == count
        assert len({digest(archive.read(name)) for name in media}) == count
    return {'pictures': count, 'distinctImageRelationships': count,
            'totalSlideRelationships': count + 1, 'uniqueMediaParts': count,
            'pythonPptxReopen': True}


def subset(parts, count):
    slide = parse(parts[SLIDE])
    tree = slide.find(P + 'cSld').find(P + 'spTree')
    pictures = tree.findall(P + 'pic')
    removed_ids = set()
    for picture in pictures[count:]:
        removed_ids.add(picture.find('.//' + A + 'blip').get(R + 'embed'))
        tree.remove(picture)
    relationships = parse(parts[RELS])
    removed_media = set()
    for relationship in list(relationships):
        if relationship.get('Id') in removed_ids:
            assert relationship.get('Type') == REL_TYPE + 'image'
            removed_media.add(media_target(relationship))
            relationships.remove(relationship)
    assert len(removed_media) == 2000 - count
    changed = {SLIDE: serialize(slide), RELS: serialize(relationships)}
    content_types = parse(parts['[Content_Types].xml'])
    removed_override = False
    for entry in list(content_types):
        if entry.get('PartName', '').lstrip('/') in removed_media:
            content_types.remove(entry)
            removed_override = True
    if removed_override:
        changed['[Content_Types].xml'] = serialize(content_types)
    result = BytesIO()
    with zipfile.ZipFile(result, 'w', compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for name in sorted(parts):
            if name in removed_media:
                continue
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o600 << 16
            archive.writestr(info, changed.get(name, parts[name]), compresslevel=9)
    return result.getvalue()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir', type=Path, required=True)
    parser.add_argument('--source', type=Path,
                        default=Path(__file__).with_name('fixtures') / 'dense2000-images.pptx')
    args = parser.parse_args()
    if pptx.__version__ != '1.0.2':
        parser.error('python-pptx 1.0.2 is required for the retained verification contract')
    output_dir = args.output_dir.resolve()
    names = [f'images-{count}.pptx' for count in COUNTS] + ['image-scale-manifest.json']
    for name in names:
        if (output_dir / name).exists():
            parser.error('refusing to overwrite ' + str(output_dir / name))
    source = args.source.resolve()
    source_data = source.read_bytes()
    if digest(source_data) != SOURCE_SHA256:
        parser.error('source does not match the pinned dense2000-images.pptx')
    source_verification = verify(source_data, 2000)
    with zipfile.ZipFile(BytesIO(source_data)) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    outputs, records = {}, []
    # Prepare and independently reopen all subsets before publishing any file.
    for count in COUNTS:
        name = f'images-{count}.pptx'
        data = subset(parts, count)
        records.append({'path': name, 'sha256': digest(data), 'bytes': len(data), **verify(data, count)})
        outputs[name] = data
    manifest = {
        'scope': 'Owned synthetic image scaling fixtures; no visual acceptance claim.',
        'source': {'path': str(source), 'sha256': SOURCE_SHA256, 'bytes': len(source_data), **source_verification},
        'generator': {'path': str(Path(__file__).resolve()), 'sha256': digest(Path(__file__).read_bytes())},
        'tools': {'python-pptx': pptx.__version__, 'python': platform.python_version(),
                  'lxml': etree.LXML_VERSION, 'zlib': zlib.ZLIB_VERSION},
        'method': 'Keep first N picture elements and their original media/relationships; fixed ZIP metadata and sorted entry order.',
        'subsets': records,
        'case2000': 'Use source.path directly; the retained original is not copied or modified.',
    }
    outputs['image-scale-manifest.json'] = (json.dumps(manifest, indent=2) + '\n').encode('utf-8')
    assert digest(source.read_bytes()) == SOURCE_SHA256
    output_dir.mkdir(parents=True, exist_ok=True)
    for name, data in outputs.items():
        # Exclusive creation protects against a concurrent file appearing
        # after the initial complete overwrite preflight.
        with (output_dir / name).open('xb') as destination:
            destination.write(data)
    print(json.dumps(manifest, indent=2))


if __name__ == '__main__':
    main()
