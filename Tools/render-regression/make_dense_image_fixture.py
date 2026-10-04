#!/usr/bin/env python3
"""Create an owned, deterministic dense image-rendering fixture (python-pptx 1.0.2)."""
import argparse
import datetime
import hashlib
import io
import json
from pathlib import Path
import struct
import zipfile
import zlib
from xml.etree import ElementTree

import pptx
from pptx import Presentation
from pptx.util import Inches


def chunk(kind, data):
    return struct.pack('>I', len(data)) + kind + data + struct.pack('>I', zlib.crc32(kind + data))


def image(number):
    # Each PNG contains a different actual RGBA pixel, rather than trailing
    # bytes that make otherwise identical image payloads appear unique.
    rgba = bytes([number & 255, (number >> 8) & 255, 128, 255])
    return (b'\x89PNG\r\n\x1a\n'
            + chunk(b'IHDR', struct.pack('>IIBBBBB', 1, 1, 8, 6, 0, 0, 0))
            + chunk(b'IDAT', zlib.compress(b'\x00' + rgba, 9))
            + chunk(b'IEND', b''))


parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('output', type=Path)
args = parser.parse_args()
output = args.output
manifest_path = output.with_suffix('.json')
if pptx.__version__ != '1.0.2':
    parser.error('python-pptx 1.0.2 required')
if output.exists() or manifest_path.exists():
    parser.error('output or manifest exists')
output.parent.mkdir(parents=True, exist_ok=True)
presentation = Presentation()
presentation.slide_width = Inches(12)
presentation.slide_height = Inches(7)
properties = presentation.core_properties
properties.author = 'Rostrum benchmark'
properties.last_modified_by = 'Rostrum benchmark'
properties.title = 'Owned dense 2000-picture benchmark'
properties.subject = '2000 distinct image relationships and unique one-pixel PNGs'
properties.created = properties.modified = datetime.datetime(2000, 1, 1)
properties.revision = 1
slide = presentation.slides.add_slide(presentation.slide_layouts[6])
for number in range(2000):
    slide.shapes.add_picture(io.BytesIO(image(number)), Inches(1), Inches(1), Inches(10), Inches(5))
raw = io.BytesIO()
presentation.save(raw)
# ZIP timestamps are transport metadata. Fix them so repeated generation
# retains a byte-for-byte fixture identity on the pinned Python environment.
with zipfile.ZipFile(raw) as source, zipfile.ZipFile(output, 'w', compression=zipfile.ZIP_DEFLATED, compresslevel=9) as destination:
    for name in sorted(source.namelist()):
        info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
        info.compress_type = zipfile.ZIP_DEFLATED
        info.create_system = 3
        info.external_attr = 0o600 << 16
        destination.writestr(info, source.read(name), compresslevel=9)

reopened = Presentation(output)
assert len(reopened.slides) == 1 and len(reopened.slides[0].shapes) == 2000
with zipfile.ZipFile(output) as archive:
    media = [name for name in archive.namelist() if name.startswith('ppt/media/')]
    assert len(media) == 2000
    assert len({hashlib.sha256(archive.read(name)).hexdigest() for name in media}) == 2000
    relationships = ElementTree.fromstring(archive.read('ppt/slides/_rels/slide1.xml.rels'))
    images = [r for r in relationships if r.attrib['Type'].endswith('/image')]
    assert len(images) == 2000 and len({r.attrib['Id'] for r in images}) == 2000
    manifest = {
        'license': 'CC0; original project-authored test geometry and image pixels',
        'scope': 'Owned synthetic rendering stress fixture; not representative image-content or visual acceptance evidence.',
        'producer': 'python-pptx ' + pptx.__version__,
        'generator': Path(__file__).name,
        'generatorSHA256': hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        'fixture': output.name,
        'fixtureSHA256': hashlib.sha256(output.read_bytes()).hexdigest(),
        'fixtureBytes': output.stat().st_size,
        'slides': 1,
        'pictures': 2000,
        'distinctImageRelationships': len(images),
        'totalSlideRelationships': len(relationships),
        'uniqueMediaParts': len(media),
        'dimensionsEMU': [presentation.slide_width, presentation.slide_height],
        'pictureFrameInches': [1, 1, 10, 5],
        'imageContent': 'Unique opaque RGBA colors in valid 1x1 PNG files, all overlapping at the same frame.',
        'verification': 'Reopened by python-pptx; ZIP counts, distinct media SHA-256 values and distinct relationship IDs asserted.',
    }
manifest_path.write_text(json.dumps(manifest, indent=2) + '\n')
print(json.dumps(manifest, indent=2))
