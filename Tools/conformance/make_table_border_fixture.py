#!/usr/bin/env python3
"""Author 42 independent shared-border cases with python-pptx 1.0.2.

Office references must be exported independently. Refuses existing outputs.
"""
import argparse
import pptx
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from pptx.oxml.xmlchemy import OxmlElement
import json,hashlib
parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('output',type=Path)
args=parser.parse_args()
if pptx.__version__!='1.0.2':parser.error('requires python-pptx 1.0.2')
manifest=args.output.with_suffix('.json')
if args.output.exists() or manifest.exists():parser.error('output exists; choose a new path')
p=Presentation();p.slide_width=Inches(12);p.slide_height=Inches(7)
cases=[]
variants=[('equal',2,2,'solid','solid',100000,100000),('first-thicker',6,1,'solid','solid',100000,100000),('second-thicker',1,6,'solid','solid',100000,100000),('first-none',2,2,'none','solid',100000,100000),('second-none',2,2,'solid','none',100000,100000),('first-absent',2,2,'absent','solid',100000,100000),('second-absent',2,2,'solid','absent',100000,100000),('first-dash',2,2,'dash','solid',100000,100000),('second-dash',2,2,'solid','dash',100000,100000),('first-alpha',2,2,'solid','solid',30000,100000),('second-alpha',2,2,'solid','solid',100000,30000)]
def edge(cell,name,color,w,kind,alpha):
 pr=cell._tc.get_or_add_tcPr()
 for e in list(pr):
  if e.tag.endswith('}'+name):pr.remove(e)
 if kind=='absent':return
 line=OxmlElement('a:'+name);line.set('w',str(int(w*12700)))
 fill=OxmlElement('a:noFill' if kind=='none' else 'a:solidFill')
 if kind!='none':
  c=OxmlElement('a:srgbClr');c.set('val',color)
  if alpha!=100000:
   a=OxmlElement('a:alpha');a.set('val',str(alpha));c.append(a)
  fill.append(c)
 line.append(fill)
 if kind=='dash':
  d=OxmlElement('a:prstDash');d.set('val','dash');line.append(d)

 order=['lnL','lnR','lnT','lnB','lnTlToBr','lnBlToTr']
 rank=order.index(name)
 index=next((i for i,e in enumerate(pr) if e.tag.split('}')[-1] not in order or order.index(e.tag.split('}')[-1])>rank),len(pr))
 pr.insert(index,line)
for direction in ['vertical','horizontal','rtl','rtl-shared','merged-vertical','merge-left','merge-top']:
 for name,w1,w2,k1,k2,a1,a2 in variants:
  if direction in ['rtl','rtl-shared','merged-vertical'] and name not in ['equal','first-thicker','second-thicker','second-none']:continue
  if direction in ['merge-left','merge-top'] and name not in ['equal','second-none','second-alpha','first-none']:continue
  s=p.slides.add_slide(p.slide_layouts[6]);rows,cols=(1,2) if direction in ['vertical','rtl','rtl-shared'] else ((2,1) if direction=='horizontal' else (2,2))
  t=s.shapes.add_table(rows,cols, Inches(1), Inches(1.5), Inches(10), Inches(4)).table
  pr=t._tbl.tblPr
  pr.attrib.clear()
  if direction in ['rtl','rtl-shared']:pr.set('rtl','1')
  pr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId').text='{2D5ABB26-0587-4C30-8999-92F81FD0307C}'
  for row in t.rows:
   for cell in row.cells:
    cell.text=''
    for en in ['lnL','lnR','lnT','lnB']:edge(cell,en,'000000',0,'none',100000)
  if direction=='merge-top':
   for col in range(cols):
    edge(t.cell(0,col),'lnB','0000FF',w1,k1,a1);edge(t.cell(1,col),'lnT','FF0000',w2,k2,a2)
   t.cell(0,0).merge(t.cell(0,1))
  elif direction=='horizontal':
   edge(t.cell(0,0),'lnB','0000FF',w1,k1,a1);edge(t.cell(1,0),'lnT','FF0000',w2,k2,a2)
  elif direction=='rtl':
   edge(t.cell(0,0),'lnL','0000FF',w1,k1,a1);edge(t.cell(0,1),'lnR','FF0000',w2,k2,a2)
  else:
   for row in range(rows):
    edge(t.cell(row,0),'lnR','0000FF',w1,k1,a1);edge(t.cell(row,1),'lnL','FF0000',w2,k2,a2)
   if direction=='merge-left':t.cell(0,0).merge(t.cell(1,0))
   if direction=='merged-vertical':t.cell(0,0).merge(t.cell(1,0));t.cell(0,1).merge(t.cell(1,1))
  cases.append({'slide':len(p.slides),'id':direction+'-'+name,'direction':direction,'first':{'color':'0000FF','widthPt':w1,'style':k1,'alpha':a1},'second':{'color':'FF0000','widthPt':w2,'style':k2,'alpha':a2}})
path=args.output;p.save(path)
manifest.write_text(json.dumps({'producer':'python-pptx 1.0.2','sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'cases':cases},indent=2)+'\n')
print(len(cases),path)
