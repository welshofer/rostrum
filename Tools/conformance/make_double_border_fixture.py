#!/usr/bin/env python3
"""Author independent double-border cases with python-pptx 1.0.2.

No Rostrum code or renderer output is used. Exports are 12 x 7 inches.
Office references must be captured separately. Refuses existing outputs.
"""
import argparse
import hashlib
import json
from pathlib import Path
import pptx
from pptx import Presentation
from pptx.util import Inches
from pptx.oxml.xmlchemy import OxmlElement

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('output', type=Path)
args = parser.parse_args()
if pptx.__version__ != '1.0.2':
    parser.error('requires python-pptx 1.0.2')
manifest = args.output.with_suffix('.json')
if args.output.exists() or manifest.exists():
    parser.error('output exists; choose a new path')
p = Presentation()
p.slide_width, p.slide_height = Inches(12), Inches(7)
cases = []
order = ['lnL', 'lnR', 'lnT', 'lnB', 'lnTlToBr', 'lnBlToTr']

def edge(cell, name, width, compound='dbl', color='0000FF', alpha=100000):
    pr = cell._tc.get_or_add_tcPr()
    for node in list(pr):
        if node.tag.split('}')[-1] == name:
            pr.remove(node)
    ln = OxmlElement('a:' + name)
    ln.set('w', str(round(width * 12700)))
    ln.set('cmpd', compound)
    fill = OxmlElement('a:solidFill' if width else 'a:noFill')
    if width:
        color_node = OxmlElement('a:srgbClr')
        color_node.set('val', color)
        if alpha != 100000:
            a = OxmlElement('a:alpha')
            a.set('val', str(alpha))
            color_node.append(a)
        fill.append(color_node)
    ln.append(fill)
    rank = order.index(name)
    index = next((i for i, node in enumerate(pr)
                  if node.tag.split('}')[-1] not in order
                  or order.index(node.tag.split('}')[-1]) > rank), len(pr))
    pr.insert(index, ln)

for direction in ['vertical', 'horizontal', 'rtl', 'merged', 'merge-left', 'merge-top', 'diagonal-down', 'diagonal-up']:
    for width in [1, 4, 9]:
        s = p.slides.add_slide(p.slide_layouts[6])
        rows, cols = ((1, 1) if direction.startswith('diagonal') else
                      (2, 1) if direction == 'horizontal' else
                      (2, 2) if direction in ['merged', 'merge-left', 'merge-top'] else (1, 2))
        t = s.shapes.add_table(rows, cols, Inches(1), Inches(1.5), Inches(10), Inches(4)).table
        pr = t._tbl.tblPr
        pr.attrib.clear()
        if direction == 'rtl':
            pr.set('rtl', '1')
        pr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId').text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'
        for row in t.rows:
            for cell in row.cells:
                cell.text = ''
                cell.fill.solid()
                cell.fill.fore_color.rgb = pptx.dml.color.RGBColor.from_string('EEE8CC')
                for name in order[:4]:
                    edge(cell, name, 0)
        if direction.startswith('diagonal'):
            edge(t.cell(0, 0), 'lnTlToBr' if direction.endswith('down') else 'lnBlToTr', width)
        elif direction in ['horizontal', 'merge-top']:
            for col in range(cols):
                edge(t.cell(0, col), 'lnB', width)
                edge(t.cell(1, col), 'lnT', 12, 'sng', 'FF0000')
            if direction == 'merge-top':
                t.cell(0, 0).merge(t.cell(0, 1))
        else:
            for row in range(rows):
                edge(t.cell(row, 0), 'lnR', width)
                edge(t.cell(row, 1), 'lnL', 12, 'sng', 'FF0000')
            if direction in ['merged', 'merge-left']:
                t.cell(0, 0).merge(t.cell(1, 0))
            if direction == 'merged':
                t.cell(0, 1).merge(t.cell(1, 1))
        cases.append({'slide': len(p.slides), 'id': f'{direction}-{width}pt', 'direction': direction,
                      'widthPt': width, 'framePx': [100, 150, 1000, 400], 'background': 'EEE8CC'})

catalog = json.loads((Path(__file__).resolve().parents[1] / 'table-style-catalog/styles.json').read_text())
for name in ['Light Style 2 - Accent 1', 'Light Style 3 - Accent 1', 'Medium Style 1 - Accent 1', 'Medium Style 3 - Accent 1', 'Dark Style 2 - Accent 1/Accent 2']:
    style = next(s for s in catalog['styles'] if s['name'] == name)
    import xml.etree.ElementTree as ET
    guid = ET.fromstring('<root xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">' + style['xml'] + '</root>')[0].attrib['styleId']
    s = p.slides.add_slide(p.slide_layouts[6])
    t = s.shapes.add_table(5, 4, Inches(1), Inches(1.5), Inches(10), Inches(4)).table
    pr = t._tbl.tblPr
    pr.attrib.clear()
    for flag in ['firstRow', 'lastRow', 'bandRow']:
        pr.set(flag, '1')
    pr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId').text = guid
    cases.append({'slide': len(p.slides), 'id': name, 'styleID': guid, 'framePx': [100, 150, 1000, 400]})
args.output.parent.mkdir(parents=True, exist_ok=True)
p.save(args.output)
manifest.write_text(json.dumps({'producer': 'python-pptx 1.0.2', 'sizePx': [1200, 700],
                               'sha256': hashlib.sha256(args.output.read_bytes()).hexdigest(), 'cases': cases}, indent=2) + '\n')
print(len(cases), args.output)
