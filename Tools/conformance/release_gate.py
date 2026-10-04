#!/usr/bin/env python3
"""Fail closed when a required independent conformance oracle is unavailable."""
import argparse, hashlib, json, subprocess, sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
p = argparse.ArgumentParser(description=__doc__)
p.add_argument('--semantic-only', action='store_true', help='development check; not a release certification')
p.add_argument('--rendered-dir', type=Path, help='candidate slide PNGs named by fixture id')
args = p.parse_args()
folder = ROOT/'Tests/RostrumTests/Fixtures/Conformance'
manifest = json.loads((folder/'manifest.json').read_text())
errors = []
try:
    import pptx
    assert pptx.__version__ == '1.0.2', 'expected python-pptx 1.0.2'
except Exception as e:
    errors.append(str(e))
for fixture in manifest['fixtures']:
    path = folder/fixture['file']
    if hashlib.sha256(path.read_bytes()).hexdigest() != fixture['sha256']:
        errors.append(f'{path.name}: fixture hash mismatch')
    try:
        deck = pptx.Presentation(path)
        table = deck.slides[0].shapes[0].table
        assert table.cell(0,0).is_merge_origin
        assert table.cell(0,0).span_height == 2 and table.cell(0,0).span_width == 2
        assert [r.font.size.pt for r in table.cell(2,1).text_frame.paragraphs[0].runs] == [18,24,14]
    except Exception as e:
        errors.append(f'{path.name}: semantic failure: {e}')
    if not args.semantic_only:
        for ref, suffix in [('powerPointReference', ''), ('notesPageReference', '-notes')]:
            if not fixture.get(ref):
                errors.append(f'{path.name}: required {ref} missing')
                continue
            reference = folder/fixture[ref]['file']
            expected = fixture[ref]['sha256']
            if not reference.is_file() or hashlib.sha256(reference.read_bytes()).hexdigest() != expected:
                errors.append(f'{path.name}: {ref} hash mismatch')
                continue
            candidate = args.rendered_dir/(fixture['id']+suffix+'.png') if args.rendered_dir else None
            if candidate is None or not candidate.is_file():
                errors.append(f'{path.name}: required {ref} candidate render missing')
            else:
                from compare_images import compare
                diff = compare(reference,candidate)
                if not diff['passed']: errors.append(f'{path.name}: {ref} visual difference: {diff}')
# Do not launch Office when mandatory semantic/reference checks already fail.
# This is still a failed release gate, not a skipped-oracle success.
if not args.semantic_only and not errors:
    for fixture in manifest['fixtures']:
        path = folder/fixture['file']
        result = subprocess.run([sys.executable,str(ROOT/'Tools/conformance/powerpoint_check.py'),str(path)])
        if result.returncode: errors.append(f'{path.name}: PowerPoint oracle failed ({result.returncode})')
print(json.dumps({'mode':'development' if args.semantic_only else 'release','errors':errors},indent=2))
sys.exit(bool(errors))
