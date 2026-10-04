#!/usr/bin/env python3
"""Author independent image-mapping cases with python-pptx 1.0.2.

Project-authored quadrant PNG, no fonts or external assets. References must be
exported separately by PowerPoint. Existing outputs are never overwritten.
"""
import argparse
import copy
import hashlib
import io
import json
from pathlib import Path
import struct
import zlib
import zipfile
import pptx
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches


def element(name, **attributes):
    node = OxmlElement(name)
    for key, value in attributes.items():
        node.set(key, str(value))
    return node


def quadrant_png():
    # 192px at 96dpi => two-inch natural size. Explicit pHYs avoids relying on
    # the producer/consumer's default image resolution.
    size = 192
    colors = [(255, 0, 0, 255), (0, 255, 0, 255),
              (0, 0, 255, 255), (255, 255, 0, 128)]
    raw = b''.join(b'\x00' + b''.join(bytes(colors[(y >= size // 2) * 2 + (x >= size // 2)])
                   for x in range(size)) for y in range(size))
    def chunk(kind, data):
        return struct.pack('>I', len(data)) + kind + data + struct.pack('>I', zlib.crc32(kind + data))
    return (b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', size, size, 8, 6, 0, 0, 0))
            + chunk(b'pHYs', struct.pack('>IIB', 3780, 3780, 1))
            + chunk(b'IDAT', zlib.compress(raw, 9)) + chunk(b'IEND', b''))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('output', type=Path)
    parser.add_argument('--without-effects', action='store_true',
                        help='explicitly disable the default shape theme shadow to isolate image mapping')
    args = parser.parse_args()
    if pptx.__version__ != '1.0.2':
        parser.error('requires python-pptx 1.0.2')
    manifest = args.output.with_suffix('.json')
    if args.output.exists() or manifest.exists():
        parser.error('output exists; choose a new path')
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(12), Inches(7)
    deck.core_properties.author = 'Rostrum conformance fixtures'
    deck.core_properties.last_modified_by = 'Rostrum conformance fixtures'
    png = quadrant_png()
    variants = [
        ('stretch', 'picture', {}, {}, {}, None),
        ('right-half-crop', 'picture', {'l': 50000}, {}, {}, None),
        ('asymmetric-crop', 'picture', {'l': 15000, 'r': 25000, 't': 10000, 'b': 20000}, {}, {}, None),
        ('negative-crop', 'picture', {'l': -50000, 'r': -50000}, {}, {}, None),
        ('destination-insets', 'picture', {}, {'l': 25000, 't': 12500, 'r': 12500, 'b': 25000}, {}, None),
        ('crop-and-destination', 'picture', {'l': -25000, 't': 25000}, {'t': 25000, 'b': 25000}, {}, None),
        ('rotate-flip-ellipse', 'picture', {'t': 50000}, {}, {'rot': 90 * 60000, 'flipH': '1', 'flipV': '1'}, None),
        ('shape-stretch', 'shape', {}, {}, {}, None),
        ('shape-tile', 'shape', {}, {}, {}, {'sx': 100000, 'sy': 100000, 'algn': 'tl'}),
        ('shape-mirror-tile', 'shape', {}, {}, {}, {'sx': 100000, 'sy': 100000, 'algn': 'ctr', 'tx': 228600, 'ty': -114300, 'flip': 'xy'}),
        ('cell-stretch', 'cell', {'l': 25000, 'b': 25000}, {}, {}, None),
        ('cell-tile', 'cell', {}, {}, {}, {'sx': 50000, 'sy': 75000, 'algn': 'br', 'flip': 'x'}),
    ]
    cases = []
    for name, kind, crop, destination, transform, tiling in variants:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        picture = slide.shapes.add_picture(io.BytesIO(png), Inches(3), Inches(2), Inches(6), Inches(3))
        fill = picture._element.blipFill
        for child in list(fill):
            if child.tag.split('}')[-1] != 'blip':
                fill.remove(child)
        if crop:
            fill.append(element('a:srcRect', **crop))
        if tiling is not None:
            fill.append(element('a:tile', **tiling))
        else:
            stretch = element('a:stretch')
            stretch.append(element('a:fillRect', **destination))
            fill.append(stretch)
        for key, value in transform.items():
            picture._element.spPr.xfrm.set(key, str(value))
        if name == 'rotate-flip-ellipse':
            picture._element.spPr.prstGeom.set('prst', 'ellipse')
        if kind != 'picture':
            image_fill = copy.deepcopy(fill)
            image_fill.tag = '{http://schemas.openxmlformats.org/drawingml/2006/main}blipFill'
            if kind == 'shape':
                shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3), Inches(2), Inches(6), Inches(3))
                props = shape._element.spPr
                for child in list(props):
                    if child.tag.split('}')[-1].endswith('Fill'):
                        props.remove(child)
                props.append(image_fill)
                line = element('a:ln')
                line.append(element('a:noFill'))
                props.append(line)
                if args.without_effects:
                    props.append(element('a:effectLst'))
                    shape._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style').find('{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef').set('idx', '0')
            else:
                table = slide.shapes.add_table(1, 1, Inches(3), Inches(2), Inches(6), Inches(3)).table
                table._tbl.tblPr.attrib.clear()
                table._tbl.tblPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId').text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'
                cell = table.cell(0, 0)
                cell.text = ''
                props = cell._tc.get_or_add_tcPr()
                for side in ('L', 'R', 'T', 'B'):
                    line = element('a:ln' + side)
                    line.append(element('a:noFill'))
                    props.append(line)
                props.append(image_fill)
            picture._element.getparent().remove(picture._element)
        cases.append({'slide': len(deck.slides), 'id': name, 'kind': kind,
                      'frameInches': [3, 2, 6, 3], 'sourceCrop': crop,
                      'destinationInsets': destination, 'transform': transform, 'tile': tiling})
    args.output.parent.mkdir(parents=True, exist_ok=True)
    buffer = io.BytesIO()
    deck.save(buffer)
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(args.output, 'w', zipfile.ZIP_DEFLATED) as output:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, (1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            output.writestr(info, source.read(name))
    record = {'producer': 'python-pptx 1.0.2', 'sourceSHA256': hashlib.sha256(args.output.read_bytes()).hexdigest(),
              'sourceImageSHA256': hashlib.sha256(png).hexdigest(), 'imagePixels': [192, 192],
              'imageDPI': [96, 96], 'slideInches': [12, 7], 'referenceStatus': 'awaiting independent Office capture',
              'shapeThemeEffectsDisabled': args.without_effects,
              'cases': cases}
    manifest.write_text(json.dumps(record, indent=2) + '\n')
    print(json.dumps({'fixture': str(args.output), 'cases': len(cases), 'sha256': record['sourceSHA256']}))


if __name__ == '__main__':
    main()
