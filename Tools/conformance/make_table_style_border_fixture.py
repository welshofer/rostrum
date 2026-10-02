#!/usr/bin/env python3
"""Author style/direct shared-edge precedence controls, python-pptx 1.0.2."""
import argparse, hashlib, json
from pathlib import Path
import xml.etree.ElementTree as ET
import pptx
from pptx import Presentation
from pptx.util import Inches
from pptx.oxml.xmlchemy import OxmlElement

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('output', type=Path)
parser.add_argument('--rtl', action='store_true')
args = parser.parse_args()
if pptx.__version__ != '1.0.2': parser.error('requires python-pptx 1.0.2')
manifest = args.output.with_suffix('.json')
if args.output.exists() or manifest.exists(): parser.error('output exists')
p = Presentation(); p.slide_width = Inches(12); p.slide_height = Inches(7)
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
catalog = json.loads((Path(__file__).resolve().parents[1] / 'table-style-catalog/styles.json').read_text())

def edge(cell, name, color):
    pr = cell._tc.get_or_add_tcPr()
    line = OxmlElement('a:' + name); line.set('w', '114300')
    fill = OxmlElement('a:solidFill' if color else 'a:noFill')
    if color:
        c = OxmlElement('a:srgbClr'); c.set('val', color); fill.append(c)
    line.append(fill)
    pr.insert(0, line)

cases = []
for region, name, boundary in [
    ('lastRow', 'Light Style 2 - Accent 1', 4),
    ('lastRow', 'Medium Style 1 - Accent 1', 4),
    ('firstRow', 'Light Style 3 - Accent 1', 1),
    ('bandRow', 'Light Style 2 - Accent 1', 2),
    ('firstCol', 'Themed Style 1 - Accent 1', 1),
    ('lastCol', 'Themed Style 1 - Accent 1', 3),
]:
    style = next(s for s in catalog['styles'] if s['name'] == name)
    guid = ET.fromstring('<r xmlns:a="' + A + '">' + style['xml'] + '</r>')[0].attrib['styleId']
    for variant in ['style-only', 'previous-solid', 'previous-none', 'current-solid', 'current-none', 'both-solid']:
        slide = p.slides.add_slide(p.slide_layouts[6])
        table = slide.shapes.add_table(5, 4, Inches(1), Inches(1.5), Inches(10), Inches(4)).table
        pr = table._tbl.tblPr; pr.attrib.clear(); pr.set(region, '1')
        if args.rtl: pr.set('rtl','1')
        pr.find('{' + A + '}tableStyleId').text = guid
        vertical = region.endswith('Col')
        for i in range(5 if vertical else 4):
            previous = table.cell(i, boundary-1) if vertical else table.cell(boundary-1, i)
            current = table.cell(i, boundary) if vertical else table.cell(boundary, i)
            if variant.startswith('previous') or variant == 'both-solid':
                edge(previous, 'lnR' if vertical else 'lnB', None if variant.endswith('none') else '0000FF')
            if variant.startswith('current') or variant == 'both-solid':
                edge(current, 'lnL' if vertical else 'lnT', None if variant.endswith('none') else 'FF0000')
        cases.append({'slide':len(p.slides),'region':region,'style':name,'variant':variant,
                      'rtl':args.rtl,'vertical':vertical,'boundary':boundary,'framePx':[100,150,1000,400]})
args.output.parent.mkdir(parents=True, exist_ok=True); p.save(args.output)
manifest.write_text(json.dumps({'producer':'python-pptx 1.0.2','sha256':hashlib.sha256(args.output.read_bytes()).hexdigest(),'sizePx':[1200,700],'cases':cases},indent=2)+'\n')
print(len(cases),args.output)
