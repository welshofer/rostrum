#!/usr/bin/env python3
"""Verify bounded notes imports with python-pptx and pinned native Office PDFs.

This checks owned fixtures and their captured output, not arbitrary notes
rendering support. It never opens Office or regenerates reference images.
"""
import argparse
import hashlib
import json
from pathlib import Path
import zipfile

import fitz
import pptx
from lxml import etree
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
FIXTURES = ROOT / 'Tests/RostrumTests/Fixtures/NotesGeometry'


def require(condition, message):
    if not condition:
        raise ValueError(message)


def geometry(slide):
    # Use the independent reader's public inherited geometry properties;
    # do not reproduce Rostrum's placeholder ancestor selection here.
    return [{'type': str(p.placeholder_format.type), 'idx': p.placeholder_format.idx,
             'rect': [p.left, p.top, p.width, p.height]}
            for p in slide.notes_slide.placeholders]


def appearance(slide):
    master = slide.notes_slide.part.notes_master
    return {
        'background': etree.tostring(master._element.find('./' + qn('p:cSld') + '/' + qn('p:bg'))),
        'style': etree.tostring(master._element.find('./' + qn('p:notesStyle'))),
        'theme': next(r.target_part.blob for r in master.part.rels.values() if r.reltype.endswith('/theme')),
        'text': slide.notes_slide.notes_text_frame.text,
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--candidates', type=Path, help='Fresh ROSTRUM_NOTES_ORACLE_OUTPUT directory')
    parser.add_argument('--type-candidates', type=Path, help='Fresh ROSTRUM_NOTES_TYPE_ORACLE_OUTPUT directory')
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    require(pptx.__version__ == '1.0.2', 'Expected python-pptx 1.0.2')
    require(fitz.VersionBind == '1.27.2.3', 'Expected PyMuPDF 1.27.2.3')
    manifest = json.loads((FIXTURES / 'manifest.json').read_text())
    decks, pdfs = {}, {}
    for name, expected in manifest['files'].items():
        path = FIXTURES / name
        require(hashlib.sha256(path.read_bytes()).hexdigest() == expected, f'{name}: reference hash mismatch')
        if path.suffix == '.pptx':
            candidate_folder = args.type_candidates if name.startswith('type-index-') else args.candidates
            if candidate_folder:
                path = candidate_folder / name
                require(hashlib.sha256(path.read_bytes()).hexdigest() == expected, f'{name}: candidate hash mismatch')
            with zipfile.ZipFile(path) as archive:
                require(archive.testzip() is None, f'{name}: ZIP CRC failure')
            decks[name] = pptx.Presentation(path)
        else:
            pdfs[name] = fitz.open(path)
    source = decks['geometry-source.pptx']
    target = decks['geometry-target-before.pptx']
    imported = decks['geometry-imported.pptx']
    require([len(d.slides) for d in [source, target, imported]] == [1, 1, 2], 'Unexpected deck slide count')
    semantic = []
    for control, ci in [(target, 0), (source, 1)]:
        require(geometry(control.slides[0]) == geometry(imported.slides[ci]), 'Inherited geometry changed')
        require(appearance(control.slides[0]) == appearance(imported.slides[ci]), 'Inherited appearance changed')
        semantic.append({'importedSlide': ci + 1, 'geometry': geometry(control.slides[0]), 'passed': True})
    extension = lambda s: etree.tostring(s.notes_slide._element.find('./' + qn('p:extLst')))
    require(extension(source.slides[0]) == extension(imported.slides[1]), 'Opaque notes XML changed')
    for kind in ['inherited', 'local']:
        before = decks[f'type-index-{kind}-source.pptx'].slides[0]
        after = decks[f'type-index-{kind}-imported.pptx'].slides[1]
        require(geometry(before) == geometry(after), f'{kind}: type-based inheritance changed')
        require(any(p['idx'] == 17 for p in geometry(after)), f'{kind}: distinct index missing')
        semantic.append({'distinctIndex': 17, 'kind': kind, 'geometry': geometry(after), 'passed': True})
    expected_pages = {'geometry-source-office.pdf': 1, 'geometry-target-before-office.pdf': 1,
                      'geometry-imported-office.pdf': 2, 'geometry-imported-reopened-office.pdf': 2}
    for name, count in expected_pages.items():
        require(len(pdfs[name]) == count, f'{name}: unexpected page count')
    native = []
    for pair in manifest['pairs']:
        left = pdfs[pair['control']][pair['controlPage']]
        right = pdfs[pair['candidate']][pair['candidatePage']]
        require(tuple(left.rect) == tuple(right.rect) == (0, 0, 612, 792), 'Expected US Letter pages')
        a = left.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        b = right.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        require(a.samples == b.samples, f'{pair}: native pixels differ')
        require(left.get_text() == right.get_text(), f'{pair}: native text differs')
        native.append(dict(pair, width=a.width, height=a.height, differingPixels=0, passed=True))
    report = {'schema': 1, 'scope': 'Pinned bounded notes geometry imports; not a universal fidelity certification',
              'pythonPptx': pptx.__version__, 'pymupdf': fitz.VersionBind,
              'freshCandidatesChecked': bool(args.candidates), 'freshTypeCandidatesChecked': bool(args.type_candidates),
              'semantic': semantic, 'native': native, 'files': manifest['files'], 'passed': True}
    args.output.write_text(json.dumps(report, indent=2) + '\n')
    print(json.dumps({'semanticPairs': len(semantic), 'nativePagePairs': len(native), 'differingPixels': 0, 'passed': True}))


if __name__ == '__main__':
    main()
