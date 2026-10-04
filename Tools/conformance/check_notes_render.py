#!/usr/bin/env python3
"""Development-only notes renderer oracle; no runtime dependency is added.

Requires python-pptx, PyMuPDF, resvg-py and Pillow. A native comparison failure
returns 1 after writing all evidence; assertion/command failures return 2.
The fixed existing visual gate is channel tolerance 16, fraction 0.005.
"""
from pathlib import Path
import argparse
import base64
import hashlib
import importlib.util
import json
import platform
import subprocess
import sys
import xml.etree.ElementTree as ET

import fitz
import pptx
from pptx.enum.shapes import PP_PLACEHOLDER
import resvg_py
from PIL import Image, ImageChops

ROOT = Path(__file__).resolve().parents[2]
FIXTURES = ROOT / 'Tests/RostrumTests/Fixtures'
CHANNEL_TOLERANCE = 16
FRACTION_TOLERANCE = 0.005
# Recorded Office exports use Letter portrait, scale-to-fit. These are print
# coordinates, not slide content bounds inferred to optimize a pixel match.
PRINT_AREA_POINTS = [17.0, 11.0, 578.0, 770.0]


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def relative(path):
    try:
        return str(path.relative_to(ROOT))
    except ValueError:
        return str(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--binary', required=True, type=Path, help='Built pptx-tool executable')
    parser.add_argument('--output-dir', required=True, type=Path, help='Candidates and native raster copies')
    parser.add_argument('--report', required=True, type=Path, help='JSON evidence path')
    parser.add_argument('--font', action='append', default=[], type=Path, help='Exact local font file; repeat for each face')
    args = parser.parse_args()
    binary, folder, report_path = args.binary.resolve(), args.output_dir.resolve(), args.report.resolve()
    fonts = [p.resolve() for p in args.font]
    folder.mkdir(parents=True, exist_ok=True)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    fontflags = sum([['--font', str(path)] for path in fonts], [])
    source_paths = sorted((FIXTURES / 'NotesGeometry').glob('*.pptx')) + [
        FIXTURES / 'Conformance' / name for name in
        ['python-tables.pptx', 'python-tables-v2.pptx', 'python-tables-v3.pptx']]
    module_hashes = {}
    for name in ['pptx', 'pptx.shapes.placeholder', 'pymupdf', 'pymupdf._mupdf', 'resvg_py.resvg_py', 'PIL._imaging']:
        spec = importlib.util.find_spec(name)
        if spec and spec.origin and Path(spec.origin).is_file():
            module_hashes[name] = {'path': spec.origin, 'sha256': sha(spec.origin)}
    report = {
        'schema': 1,
        'command': [sys.executable, str(Path(__file__).resolve())] + sys.argv[1:],
        'revision': subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=ROOT, text=True).strip(),
        'engines': {'python': platform.python_version(), 'python-pptx': pptx.__version__,
                    'resvg_py': resvg_py.__version__, 'PyMuPDF': fitz.VersionBind, 'Pillow': Image.__version__},
        'tools': {'scriptSHA256': sha(__file__), 'binarySHA256': sha(binary),
                  'pythonExecutableSHA256': sha(sys.executable), 'modules': module_hashes},
        'implementationSHA256': {relative(p): sha(p) for p in [
            ROOT / 'Sources/Rostrum/Presentation/NotesPageRenderer.swift',
            ROOT / 'Sources/Rostrum/Presentation/SVGRenderer.swift',
            ROOT / 'Sources/Rostrum/Presentation/RichTextLayout.swift']},
        'allLibrarySourcesSHA256': hashlib.sha256(json.dumps({relative(p): sha(p)
            for p in sorted((ROOT / 'Sources/Rostrum').rglob('*.swift'))}, sort_keys=True).encode()).hexdigest(),
        'sourceSHA256': {relative(p): sha(p) for p in source_paths},
        'fontSHA256': {str(p): sha(p) for p in fonts},
        'nativePDFSHA256': {},
        'apiSizeRendering': {'preservesSourceNotesSize': True, 'decks': []},
        'sourceImportPairs': [],
        'printProfileNormalization': {
            'paperPoints': [612, 792], 'notesAreaPoints': PRINT_AREA_POINTS,
            'explanation': 'Explicit recorded Office print profile applied only to comparison copies; API SVG keeps p:notesSz. No per-page fit to content or reference changes.',
            'channelTolerance': CHANNEL_TOLERANCE, 'maximumDifferenceFraction': FRACTION_TOLERANCE,
            'comparisons': [],
        },
    }
    namespace = {'s': 'http://www.w3.org/2000/svg'}
    for path in source_paths:
        out = folder / path.stem
        command = [str(binary), 'render', str(path), str(out), '--notes'] + fontflags
        result = subprocess.run(command, text=True, capture_output=True, cwd=ROOT)
        if result.returncode:
            raise RuntimeError(f'{command!r} exited {result.returncode}: {result.stderr}')
        oracle = pptx.Presentation(path)
        notes_size = oracle._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}notesSz')
        expected_size = [int(notes_size.get('cx')), int(notes_size.get('cy'))]
        records = []
        for index, slide in enumerate(oracle.slides):
            if not slide.has_notes_slide:
                continue
            svg = out / f'notes-{index + 1:02d}.svg'
            root = ET.fromstring(svg.read_text())
            viewbox = [int(value) for value in root.attrib['viewBox'].split()]
            assert viewbox == [0, 0] + expected_size, (path, viewbox, expected_size)
            image = root.find('s:image', namespace)
            inherited = next((shape for shape in slide.notes_slide.placeholders
                              if shape.placeholder_format.type == PP_PLACEHOLDER.SLIDE_IMAGE), None)
            expected = [inherited.left, inherited.top, inherited.width, inherited.height] if inherited is not None else None
            actual = [int(image.attrib[key]) for key in ['x', 'y', 'width', 'height']] if image is not None else None
            assert expected == actual, (path, index, expected, actual)
            png = svg.with_suffix('.png')
            png.write_bytes(resvg_py.svg_to_bytes(svg_path=str(svg), width=750))
            records.append({'slide': index + 1, 'svg': str(svg), 'svgSHA256': sha(svg), 'pngSHA256': sha(png),
                            'notesSizeEMU': expected_size, 'pythonPptxImageFrame': expected, 'renderImageFrame': actual})
        report['apiSizeRendering']['decks'].append({'input': relative(path), 'exitCode': result.returncode,
                                                   'diagnostics': result.stderr.splitlines(), 'pages': records})
    pairs = [(folder / f'{kind}-source/notes-01.svg', folder / f'{kind}-imported/notes-02.svg')
             for kind in ['geometry', 'type-index-inherited', 'type-index-local']]
    pairs.append((folder / 'geometry-target-before/notes-01.svg', folder / 'geometry-imported/notes-01.svg'))
    for source, imported in pairs:
        assert source.read_bytes() == imported.read_bytes(), (source, imported)
        report['sourceImportPairs'].append({'source': str(source), 'imported': str(imported),
                                           'byteIdenticalSVG': True, 'sharedSHA256': sha(source)})
    native_cases = [(FIXTURES / 'NotesGeometry' / f'{name}-office.pdf', name)
                    for name in ['geometry-source', 'geometry-target-before', 'geometry-imported']]
    native_cases += [(FIXTURES / 'Conformance' / f'{name}-notes-office.pdf', name)
                     for name in ['python-tables', 'python-tables-v2', 'python-tables-v3']]
    comparisons = report['printProfileNormalization']['comparisons']
    for pdf, name in native_cases:
        report['nativePDFSHA256'][relative(pdf)] = sha(pdf)
        with fitz.open(pdf) as native:
            for index, page in enumerate(native):
                svg = folder / name / f'notes-{index + 1:02d}.svg'
                paper = page.rect
                assert (paper.width, paper.height) == (612, 792), (pdf, paper)
                x, y, width, height = PRINT_AREA_POINTS
                content = base64.b64encode(svg.read_bytes()).decode()
                wrapper = (f'<svg xmlns="http://www.w3.org/2000/svg" width="1200" height="1553" viewBox="0 0 {paper.width} {paper.height}">'
                           f'<rect width="100%" height="100%" fill="white"/>'
                           f'<image x="{x}" y="{y}" width="{width}" height="{height}" preserveAspectRatio="none" '
                           f'href="data:image/svg+xml;base64,{content}"/></svg>')
                candidate = svg.with_name(svg.stem + '-print-candidate.png')
                candidate.write_bytes(resvg_py.svg_to_bytes(svg_string=wrapper, width=1200))
                reference = svg.with_name(svg.stem + '-native.png')
                page.get_pixmap(matrix=fitz.Matrix(1200 / paper.width, 1200 / paper.width), alpha=False).save(reference)
                a, b = Image.open(candidate).convert('RGB'), Image.open(reference).convert('RGB')
                assert a.size == b.size, (a.size, b.size)
                delta = ImageChops.difference(a, b)
                channels = delta.split()
                maximum = ImageChops.lighter(ImageChops.lighter(channels[0], channels[1]), channels[2])
                difference_count = sum(maximum.histogram()[CHANNEL_TOLERANCE + 1:])
                fraction = difference_count / (a.width * a.height)
                comparisons.append({'deck': name, 'page': index + 1, 'candidate': str(candidate),
                                    'nativeRaster': str(reference), 'referencePDF': relative(pdf),
                                    'candidateSHA256': sha(candidate), 'nativeRasterSHA256': sha(reference),
                                    'pixelSize': list(a.size), 'differingPixelFraction': fraction,
                                    'passed': fraction <= FRACTION_TOLERANCE})
    report['printProfileNormalization']['passed'] = all(row['passed'] for row in comparisons)
    report['passed'] = report['printProfileNormalization']['passed']
    report_path.write_text(json.dumps(report, indent=2) + '\n')
    decks = report['apiSizeRendering']['decks']
    print(f'{len(decks)} decks reopened; {sum(len(d["pages"]) for d in decks)} source-size notes candidates; '
          '4 exact source/import SVG pairs; python-pptx placeholder geometry agrees.')
    for row in comparisons:
        print(f'{row["deck"]} page {row["page"]}: {row["differingPixelFraction"]:.9f} '
              f'({"PASS" if row["passed"] else "FAIL"}, fixed limit {FRACTION_TOLERANCE})')
    print(f'Evidence: {report_path}')
    return 0 if report['passed'] else 1


if __name__ == '__main__':
    try:
        sys.exit(main())
    except (AssertionError, RuntimeError, OSError, ValueError) as error:
        print(f'notes render oracle failed: {error}', file=sys.stderr)
        sys.exit(2)
