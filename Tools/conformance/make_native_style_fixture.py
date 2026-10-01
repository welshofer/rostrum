#!/usr/bin/env python3
"""Author the original 74-slide native-style probe with pinned python-pptx 1.0.2.

Only native GUIDs/names are used, never catalog definition XML. Existing files
are never replaced. PowerPoint PNG references require an independent export.
"""
import argparse
import json
from pathlib import Path
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('output', type=Path)
args = parser.parse_args()
if pptx.__version__ != '1.0.2': parser.error('requires python-pptx 1.0.2')
if args.output.exists(): parser.error('output exists; choose a new path')
root = Path(__file__).resolve().parents[2]
styles = json.loads((root/'Tools/table-style-catalog/styles.json').read_text())['styles']
deck = Presentation()
deck.slide_width, deck.slide_height = Inches(12), Inches(7)
deck.core_properties.author = deck.core_properties.last_modified_by = 'Rostrum conformance'
for style in styles:
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(.5), Inches(.2), Inches(11), Inches(.55))
    title.text_frame.text = style['name']
    for run in title.text_frame.paragraphs[0].runs:
        run.font.name, run.font.size = 'Arial', Pt(22)
    table = slide.shapes.add_table(5, 4, Inches(.5), Inches(1), Inches(11), Inches(4)).table
    table.first_row, table.horz_banding = True, True
    table._tbl.tblPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId').text = style['id']
    for row in range(5):
        for column in range(4):
            cell = table.cell(row, column)
            cell.text = f'{row+1}.{column+1}'
            for run in cell.text_frame.paragraphs[0].runs:
                run.font.name, run.font.size = 'Arial', Pt(18)
    slide.notes_slide.notes_text_frame.text = 'Rostrum native table-style verification. Original generated fixture.'
deck.save(args.output)
