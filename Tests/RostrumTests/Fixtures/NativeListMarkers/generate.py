"""Independent list-marker discriminator matrix, never Rostrum-derived geometry."""
from pathlib import Path
from io import BytesIO
import hashlib,json,struct,zipfile,math
from lxml import etree
from fontTools.ttLib import TTFont
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import MSO_AUTO_SIZE
OUT=Path(__file__).resolve().parent
ROOT=next(p for p in OUT.parents if (p/'Package.swift').exists())
FONTS=ROOT/'Tests/RostrumTests/Fixtures/NativeGlyphPlacement/eligibility/fonts'
FACES={
 'regular':dict(path=FONTS/'DejaVuSans.ttf',family='DejaVu Sans',style='Book',slot='regular',bold=False,italic=False),
 'serif':dict(path=FONTS/'DejaVuSerif.ttf',family='DejaVu Serif',style='Book',slot='regular',bold=False,italic=False),
 'bold':dict(path=FONTS/'DejaVuSans-Bold.ttf',family='DejaVu Sans',style='Bold',slot='bold',bold=True,italic=False),
}
A='http://schemas.openxmlformats.org/drawingml/2006/main';P='http://schemas.openxmlformats.org/presentationml/2006/main';R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
def elem(name,**attrs):return etree.Element('{'+A+'}'+name,**{k:str(v) for k,v in attrs.items()})
def run(text,size):return dict(kind='run',text=text,size=size)
cases=[]
def add(name,size,**kw):
 c=dict(name=name,size=size,marker='•',markerType='character',sizeChoice='text',fontChoice='text',face='regular',fontScale=100,marL=24,indent=-18,wrap=False,width=290,height=160,nodes=[run('Agjp BBBB Z',size)])
 c.update(kw);cases.append(c)
add('follow-text-integer16',16)
add('follow-text-authored14p5',14.5)
add('follow-text-scaled29x50',29,fontScale=50)
add('oversized150-two-lines',18,sizeChoice='percent',markerSize=150,nodes=[run('Agjp BBBB',18),dict(kind='break',size=18),run('Agjp Z',18)])
add('percent75-authored14p5',14.5,sizeChoice='percent',markerSize=75,competingPaintSizes=[11,10])
add('percent90-scaled29x50',29,fontScale=50,sizeChoice='percent',markerSize=90,competingPaintSizes=[13,14])
add('points14p5-body18',18,sizeChoice='points',markerSize=14.5)
add('points14p5-scaled20x72p5',20,fontScale=72.5,sizeChoice='points',markerSize=14.5,competingPaintSizes=[11,10,14])
add('inherited-percent75-serif',14.5,inherited=True,sizeChoice='inherited',fontChoice='inherited')
add('local-follow-overrides-inherited',14.5,inherited=True)
add('follow-first-run20-default14p5',14.5,nodes=[run('Agjp ',20),run('BBBB Z',14.5)])
add('follow-actual-bold14p5',14.5,face='bold')
add('number12-integer18',18,markerType='number',marker='12.',startAt=12)
add('number12-scaled29x50',29,fontScale=50,markerType='number',marker='12.',startAt=12)
add('number12-explicit-serif',18,markerType='number',marker='12.',startAt=12,fontChoice='serif')
add('character-explicit-serif14p5',14.5,fontChoice='serif')
for indent in [-18,-6]:
 add('wide-number-hang'+str(-indent),18,markerType='number',marker='12345.',startAt=12345,indent=indent,width=130,height=190 if indent==-6 else 160,wrap=True,nodes=[run('BBBB BBBB BBBB BBBB Z',18)])
assert len(cases)==18

def properties(name,size,face='regular'):
 f=FACES[face];node=elem(name,sz=round(size*100),b=int(f['bold']),i=int(f['italic']),lang='en-US',kern=0)
 fill=elem('solidFill');fill.append(elem('srgbClr',val='000000'));node.append(fill);node.append(elem('latin',typeface=f['family']));return node

def bullet_choices(pr,c,include_marker=True):
 choice=c['sizeChoice']
 if choice=='text':pr.append(elem('buSzTx'))
 elif choice=='percent':pr.append(elem('buSzPct',val=round(c['markerSize']*1000)))
 elif choice=='points':pr.append(elem('buSzPts',val=round(c['markerSize']*100)))
 if c['fontChoice']=='text':pr.append(elem('buFontTx'))
 elif c['fontChoice']=='serif':pr.append(elem('buFont',typeface='DejaVu Serif'))
 if include_marker:
  pr.append(elem('buAutoNum',type='arabicPeriod',startAt=c['startAt']) if c['markerType']=='number' else elem('buChar',char=c['marker']))

deck=Presentation();deck.slide_width=Pt(720);deck.slide_height=Pt(720)
# The two inheritance controls resolve this actual master otherStyle choice.
master=deck.slide_masters[0]._element
other=master.find('{'+P+'}txStyles').find('{'+P+'}otherStyle')
level=other.find('{'+A+'}lvl1pPr')
for child in list(level):level.remove(child)
level.append(elem('buSzPct',val=75000));level.append(elem('buFont',typeface='DejaVu Serif'));level.append(elem('buChar',char='•'));level.append(properties('defRPr',14.5))
for index,c in enumerate(cases):
 if index%6==0:slide=deck.slides.add_slide(deck.slide_layouts[6])
 x=30+350*(index%2);y=50+220*((index%6)//2)
 label=slide.shapes.add_textbox(Pt(x),Pt(y-23),Pt(310),Pt(18));label.text=c['name'];label.text_frame.paragraphs[0].runs[0].font.size=Pt(10)
 # Labels explicitly suppress inherited markers and are excluded by frame.
 label.text_frame.paragraphs[0]._p.get_or_add_pPr().append(elem('buNone'))
 shape=slide.shapes.add_textbox(Pt(x),Pt(y),Pt(c['width']),Pt(c['height']));shape.name=c['name'];tf=shape.text_frame
 tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.word_wrap=c['wrap'];tf.auto_size=MSO_AUTO_SIZE.NONE
 body=tf._txBody;bp=body.find('{'+A+'}bodyPr');bp.set('anchor','t')
 if c['fontScale']!=100:
  for child in list(bp):
   if etree.QName(child).localname in ['normAutofit','noAutofit','spAutoFit']:bp.remove(child)
  bp.append(elem('normAutofit',fontScale=round(c['fontScale']*1000),lnSpcReduction=0))
 for child in list(body):
  if etree.QName(child).localname in ['p','lstStyle']:body.remove(child)
 body.append(elem('lstStyle'));p=elem('p');pr=elem('pPr',algn='l',lvl=0,marL=round(c['marL']*12700),indent=round(c['indent']*12700))
 bullet_choices(pr,c,include_marker=not c.get('inherited',False))
 # Inherited case deliberately has no local size or font; body run properties
 # still pin actual face and size independently of marker choice inheritance.
 if not c.get('inherited',False):pr.append(properties('defRPr',c['size'],c['face']))
 p.append(pr)
 for token in c['nodes']:
  node=elem('br' if token['kind']=='break' else 'r');node.append(properties('rPr',token['size'],c['face']))
  if token['kind']=='run':t=elem('t');t.text=token['text'];node.append(t)
  p.append(node)
 body.append(p);c.update(page=index//6,x=x,y=y,sourceTextBodyXML=etree.tostring(body).decode(),expectedVisibleText=c['marker']+''.join(n.get('text','') for n in c['nodes']).replace(' ',''))
# Independently embed actual licensed faces in their correct style slots.
buffer=BytesIO();deck.save(buffer)
with zipfile.ZipFile(buffer) as source:parts={n:source.read(n) for n in source.namelist()}
pres=etree.fromstring(parts['ppt/presentation.xml']);pres.set('embedTrueTypeFonts','1');pres.set('saveSubsetFonts','0')
fonts=etree.Element('{'+P+'}embeddedFontLst');entries={}
rels=etree.fromstring(parts['ppt/_rels/presentation.xml.rels']);face_manifest={}
for key,face in FACES.items():
    font_bytes=face['path'].read_bytes();font=TTFont(face['path']);header=bytearray(82)
    struct.pack_into('<IIII',header,0,0,len(font_bytes),0x00020001,0);header[26]=1;header[27]=1 if face['italic'] else 0
    struct.pack_into('<IHH',header,28,font['OS/2'].usWeightClass,font['OS/2'].fsType,0x504c)
    names=bytearray()
    for value in [face['family'],face['style'],font['name'].getDebugName(5),face['family']]:
        encoded=value.encode('utf-16le');names+=struct.pack('<H',len(encoded))+encoded+b'\0\0'
    eot=header+names+b'\0\0'+font_bytes;struct.pack_into('<I',eot,0,len(eot))
    if face['family'] not in entries:
        entry=etree.SubElement(fonts,'{'+P+'}embeddedFont');etree.SubElement(entry,'{'+P+'}font').set('typeface',face['family']);entries[face['family']]=entry
    relationship='rIdPaint'+key.title()
    etree.SubElement(entries[face['family']],'{'+P+'}'+face['slot']).set('{'+R+'}id',relationship)
    etree.SubElement(rels,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',Id=relationship,Type=R+'/font',Target='fonts/'+key+'.fntdata')
    parts['ppt/fonts/'+key+'.fntdata']=eot
    (OUT/'fonts').mkdir(exist_ok=True);(OUT/'fonts'/face['path'].name).write_bytes(font_bytes)
    face_manifest[key]=dict(family=face['family'],bold=face['bold'],italic=face['italic'],file='fonts/'+face['path'].name,sha256=hashlib.sha256(font_bytes).hexdigest(),originalPath=str(face['path']),embeddingPermissions=font['OS/2'].fsType)
pres.insert(list(pres).index(pres.find('{'+P+'}defaultTextStyle')),fonts)
def xml(node):return etree.tostring(node,xml_declaration=True,encoding='UTF-8',standalone=True)
parts['ppt/presentation.xml']=xml(pres);parts['ppt/_rels/presentation.xml.rels']=xml(rels)
ct=etree.fromstring(parts['[Content_Types].xml']);etree.SubElement(ct,'{http://schemas.openxmlformats.org/package/2006/content-types}Default',Extension='fntdata',ContentType='application/x-fontdata');parts['[Content_Types].xml']=xml(ct)
output=OUT/'native-list-markers-v2.pptx'
with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as target:
    for name,data in sorted(parts.items()):target.writestr(name,data)
(OUT/'cases.json').write_text(json.dumps(cases,indent=2)+'\n')
manifest=dict(source=output.name,sourceSHA256=hashlib.sha256(output.read_bytes()).hexdigest(),slideCount=3,caseCount=len(cases),faces=face_manifest,scope='Independent list-marker paint, measurement, inheritance and hanging-indent controls. No native outcome assumed.')
(OUT/'manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
print(output);print(manifest['sourceSHA256']);print('3 slides,18 cases')
