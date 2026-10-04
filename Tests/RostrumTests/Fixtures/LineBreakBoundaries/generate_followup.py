"""Independent native probes: fontTools selects widths; python-pptx authors OOXML.
No Rostrum measurements/serialization are used. Candidate grid numbers choose
inputs only; native PDF line contents remain the independent expected outputs.
"""
from pathlib import Path
from io import BytesIO
import hashlib,json,math,struct,zipfile
from lxml import etree
from fontTools.ttLib import TTFont
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN

out=Path(__file__).parent
root=out.parents[3]
paths={
 'Arial':Path('/System/Library/Fonts/Supplemental/Arial.ttf'),
 'Arial Bold':Path('/System/Library/Fonts/Supplemental/Arial Bold.ttf'),
 'DejaVu Sans':root/'Lectern/Sources/LecternCore/Resources/LibraryLab/DejaVuSans.ttf',
}
fonts={k:TTFont(v) for k,v in paths.items()}
def run(text,face='Arial',size=18,tracking=0,kern=False):
 return dict(text=text,face=face,size=size,tracking=tracking,kern=kern)
known='SupercalifragilisticexpialidociousSupe'
groups=[]
for kern in [None,0,100,1200,1800,1801,400000]:
 for spacing in [None,0,.25]:
  r=run('AVATARAVATAR',kern=True)
  r['kernValue']=kern;r['spacingValue']=spacing
  groups.append((f'kern-{kern}-spacing-{spacing}',[r],'A',[200]))
for kern,spacing in [(None,None),(0,None),(1200,None),(1200,0)]:
 r=run('AVATARAVATAR',kern=True);r['kernValue']=kern;r['spacingValue']=spacing
 groups.append((f'kern-boundary-{kern}-spacing-{spacing}',[r],'A',[133.30,133.33,143.99,144.01]))
groups += [
 ('tracking-positive',[dict(run('B'*12,size=10,tracking=.07),kernValue=400000,spacingValue=.07)],'Z',
  [80.30,80.33,80.34,80.35,80.36,80.374,80.376,80.40,80.49,80.51,80.74,80.76,80.84,80.86]),
 ('tracking-negative',[dict(run('B'*12,size=10,tracking=-.07),kernValue=400000,spacingValue=-.07)],'Z',
  [78.64,78.66,78.68,78.74,78.76,78.99,79.01]),
 ('odd-glyph-count',[dict(run('B'*11,size=10),kernValue=400000,spacingValue=None)],'Z',
  [72.85,72.865,72.875,72.885,72.90,72.99,73.01]),
]
def measure(r,grid=False):
 f=fonts[r['face']]; cm=f.getBestCmap(); upem=f['head'].unitsPerEm
 widths=[f['hmtx'].metrics[cm[ord(c)]][0]*r['size']/upem for c in r['text']]
 if grid: widths=[math.floor(w*8+.5)/8 for w in widths]
 total=sum(widths)+len(widths)*r['tracking']
 if r['kern'] and 'kern' in f:
  pairs={}
  for table in f['kern'].kernTables: pairs.update(table.kernTable)
  total+=sum(pairs.get((cm[ord(a)],cm[ord(b)]),0) for a,b in zip(r['text'],r['text'][1:]))*r['size']/upem
 return total
p=Presentation();p.slide_width=Pt(720);p.slide_height=Pt(540)
cases=[]
for name,runs,suffix,widths in groups:
 raw=sum(measure(r) for r in runs); grid=sum(measure(r,True) for r in runs)
 for probe,width in enumerate(widths):
  i=len(cases)
  if i%6==0: slide=p.slides.add_slide(p.slide_layouts[6])
  x,y=30+(i%2)*350,20+((i%6)//2)*170
  label=slide.shapes.add_textbox(Pt(x),Pt(y),Pt(320),Pt(20))
  label.text=f'{name}-{probe}: {width:.5f}pt'
  label.text_frame.paragraphs[0].runs[0].font.size=Pt(10)
  box=slide.shapes.add_textbox(Pt(x),Pt(y+25),Pt(width),Pt(125));box.name=f'{name}-{probe}'
  tf=box.text_frame;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
  tf.word_wrap=True;tf.auto_size=MSO_AUTO_SIZE.NONE
  para=tf.paragraphs[0];para.alignment=PP_ALIGN.LEFT;para.space_before=para.space_after=Pt(0)
  authored=[dict(r) for r in runs];authored[-1]['text']+=suffix
  for r in authored:
   rr=para.add_run();rr.text=r['text'];rr.font.name='Arial' if r['face']=='Arial Bold' else r['face']
   rr.font.size=Pt(r['size']);rr.font.bold=r['face']=='Arial Bold'
   props=rr._r.get_or_add_rPr();props.set('lang','en-US');
   if r.get('kernValue') is not None: props.set('kern',str(r['kernValue']))
   if r.get('spacingValue') is not None: props.set('spc',str(round(r['spacingValue']*100)))
  cases.append(dict(name=box.name,group=name,page=i//6,x=x,y=y+25,width=box.width/12700,height=125,
                    runs=authored,prefix=''.join(r['text'] for r in runs),rawPrefixWidth=raw,candidateGridPrefixWidth=grid))
buf=BytesIO();p.save(buf)
# Independent EOT v2.1 construction per https://www.w3.org/submissions/EOT/ §3.2.
# The redistributable DejaVu font is embedded to pin the native face on macOS.
font=paths['DejaVu Sans'].read_bytes();ft=fonts['DejaVu Sans']
header=bytearray(82)
struct.pack_into('<IIII',header,0,0,len(font),0x00020001,0)
header[26]=1
struct.pack_into('<IHH',header,28,400,ft['OS/2'].fsType,0x504c)
names=bytearray()
for name in ['DejaVu Sans','Book','Version 2.37','DejaVu Sans']:
 encoded=name.encode('utf-16le');names+=struct.pack('<H',len(encoded))+encoded+b'\0\0'
eot=header+names+b'\0\0'+font;struct.pack_into('<I',eot,0,len(eot))
NSP='http://schemas.openxmlformats.org/presentationml/2006/main'; NSR='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
with zipfile.ZipFile(buf) as source: parts={n:source.read(n) for n in source.namelist()}
pres=etree.fromstring(parts['ppt/presentation.xml']);pres.set('embedTrueTypeFonts','1');pres.set('saveSubsetFonts','0')
el=etree.Element('{'+NSP+'}embeddedFontLst');ef=etree.SubElement(el,'{'+NSP+'}embeddedFont')
etree.SubElement(ef,'{'+NSP+'}font').set('typeface','DejaVu Sans');etree.SubElement(ef,'{'+NSP+'}regular').set('{'+NSR+'}id','rIdBoundaryFont')
default=pres.find('{'+NSP+'}defaultTextStyle');pres.insert(list(pres).index(default),el)
parts['ppt/presentation.xml']=etree.tostring(pres,xml_declaration=True,encoding='UTF-8',standalone=True)
rels=etree.fromstring(parts['ppt/_rels/presentation.xml.rels'])
rel=etree.SubElement(rels,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',Id='rIdBoundaryFont',Type=NSR+'/font',Target='fonts/boundary.fntdata')
parts['ppt/_rels/presentation.xml.rels']=etree.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
ct=etree.fromstring(parts['[Content_Types].xml']);etree.SubElement(ct,'{http://schemas.openxmlformats.org/package/2006/content-types}Default',Extension='fntdata',ContentType='application/x-fontdata')
parts['[Content_Types].xml']=etree.tostring(ct,xml_declaration=True,encoding='UTF-8',standalone=True);parts['ppt/fonts/boundary.fntdata']=eot
with zipfile.ZipFile(out/'line-break-followup.pptx','w',zipfile.ZIP_DEFLATED) as target:
 for name,data in sorted(parts.items()): target.writestr(name,data)
(out/'cases-followup.json').write_text(json.dumps(cases,indent=2)+'\n')
manifest=dict(generator='python-pptx 1.0.2 and fontTools; independent OOXML/EOT construction',candidate='Follow-up isolates text-box width thresholds, tracking and independent kern/spacing attributes; native outcomes determine support.',fonts=[dict(name=n,path=str(v),sha256=hashlib.sha256(v.read_bytes()).hexdigest()) for n,v in paths.items()])
(out/'input-manifest-followup.json').write_text(json.dumps(manifest,indent=2)+'\n')
print(len(cases),'cases;',len(p.slides),'slides;',out/'line-break-followup.pptx')
