from pathlib import Path
import hashlib,json,zipfile,struct
from lxml import etree
from pptx import Presentation
from fontTools.ttLib import TTFont
R=Path(__file__).resolve().parent
m=json.loads((R/'manifest.json').read_text());c=json.loads((R/'cases.json').read_text());p=R/m['source']
assert hashlib.sha256(p.read_bytes()).hexdigest()==m['sourceSHA256']
assert len(c)==24 and len({v['name'] for v in c})==24
prs=Presentation(p);assert len(prs.slides)==4
A='{http://schemas.openxmlformats.org/drawingml/2006/main}'
P='{http://schemas.openxmlformats.org/presentationml/2006/main}'
with zipfile.ZipFile(p) as z:
 assert z.testzip() is None
 for n in z.namelist():
  if n.endswith(('.xml','.rels')):etree.fromstring(z.read(n))
 for case in c:
  shapes=[s for s in prs.slides[case['page']].shapes if s.name==case['name']];assert len(shapes)==1
  shape=shapes[0];tf=shape.table.cell(0,0).text_frame if case.get('table') else shape.text_frame
  assert shape.width==round(case['width']*12700)
  body=tf._txBody
  assert body.find(A+'p').find(A+'pPr').get('algn')==case['alignment']
  assert list(body)[0].tag==A+'bodyPr' and list(body)[1].tag==A+'lstStyle'
  expected=''.join(n.get('text','') if n['kind']=='run' else '\v' for q in case['paragraphs'] for n in q['nodes'])
  assert tf.text==expected,(case['name'],repr(tf.text),repr(expected))
  assert not shape._element.xpath('.//p:timing')
 for key,face in m['faces'].items():
  b=(R/face['file']).read_bytes();assert hashlib.sha256(b).hexdigest()==face['sha256']
  e=z.read('ppt/fonts/'+key+'.fntdata');size=struct.unpack_from('<I',e,4)[0];assert e[-size:]==b
  assert TTFont(R/face['file'])['OS/2'].fsType==0
result=dict(sourceSHA256=m['sourceSHA256'],cases=24,slides=4,visibleScalars=sum(len([v for v in n.get('text','') if v!=' ']) for x in c for p in x['paragraphs'] for n in p['nodes']),packageXMLAndRelationshipsParsed=True,sourceBodyTextWidthsAlignmentAndChildOrderMatch=True,embeddedSFNTExact=True,productionBaseline='05613d0',nativeAcceptance='Accepted PDF separately pinned by capture-receipt.json; glyph proof in native-alignment-metrics.json')
(R/'audit.json').write_text(json.dumps(result,indent=2)+'\n');print(json.dumps(result))
