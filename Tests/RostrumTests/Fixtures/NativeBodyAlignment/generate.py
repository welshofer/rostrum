"""Independent native glyph-ink, placement and wrap discriminators; no Rostrum geometry."""
from pathlib import Path
from io import BytesIO
import hashlib, json, struct, zipfile, math, shutil
from lxml import etree
from fontTools.ttLib import TTFont
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import MSO_AUTO_SIZE

OUT=Path(__file__).resolve().parent
ROOT=next(p for p in OUT.parents if (p/'Package.swift').exists())
FONT=ROOT/'Tests/RostrumTests/Fixtures/Typography/DejaVuSans.ttf'
BUNDLED=ROOT/'Tests/RostrumTests/Fixtures/NativeListMarkers/fonts'
FACES={
 'regular':dict(path=BUNDLED/'DejaVuSans.ttf',family='DejaVu Sans',style='Book',slot='regular',bold=False,italic=False),
 'serif':dict(path=BUNDLED/'DejaVuSerif.ttf',family='DejaVu Serif',style='Book',slot='regular',bold=False,italic=False),
 'bold':dict(path=BUNDLED/'DejaVuSans-Bold.ttf',family='DejaVu Sans',style='Bold',slot='bold',bold=True,italic=False),
}
A='http://schemas.openxmlformats.org/drawingml/2006/main'
P='http://schemas.openxmlformats.org/presentationml/2006/main'
R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
def elem(name,**attrs):return etree.Element('{'+A+'}'+name,**{k:str(v) for k,v in attrs.items()})
def run(text,size=12):return dict(kind='run',text=text,size=size)
def br(size=None):return dict(kind='break',size=size)
def paragraph(nodes,default=None,end=None,spacing=None):return dict(nodes=nodes,default=default,end=end,spacing=spacing)
def pair(size):return [paragraph([run('A'),br(size),run('B')])]
def consecutive(first,second,default=None):return [paragraph([run('A'),br(first),br(second),run('B')],default=default)]
def trailing(end):return [paragraph([run('A'),br(36)],end=end),paragraph([run('Z')])]
def empty(end):return [paragraph([run('A')]),paragraph([],end=end),paragraph([run('Z')])]
cases=[]
def add(name, runs, **kwargs):
    cases.append(dict(name=name,paragraphs=[paragraph(runs)],**kwargs))
def paired(name, nodes, **kwargs):
    for alignment in ['ctr','r']:
        add(name+'-'+alignment,nodes,alignment=alignment,**kwargs)
for width in [110.99,111.01]:
    paired(f'fractional-width-{width}',[run('Agjp',14.5)],width=width,wrap=False)
for left,right in [(0,0),(7.2,3.6)]:
    paired(f'multiline-trailing-space-insets-{left}-{right}',[run('Agjp  ',14.5),br(14.5),run('BBBB',14.5)],width=111.01,wrap=False,insets=[left,right,0,0])
# Independent hmtx calculation, before native capture: twelve B advances are
# exactly 111 pt at the already calibrated unscaled 13.51 pt eighth-point grid.
f=TTFont(FACES['regular']['path']);u=f['head'].unitsPerEm
advance=f['hmtx'][f.getBestCmap()[ord('B')]][0]
quantized=math.floor(advance*13.51/u*8+.5)/8
assert quantized==9.25 and 12*quantized==111
for width in [110.99,111.01]:
    paired(f'wrap-prefix111-width-{width}',[run('BBBBBBBBBBBBZ',13.51)],width=width,wrap=True,
           discriminator=dict(prefix='BBBBBBBBBBBB',sourceAdvance=advance,unitsPerEm=u,
                              authoredSize=13.51,perScalarGridAdvance=quantized,prefixGridAdvance=12*quantized,
                              rawCapacity=width,eighthFloorCapacity=math.floor(width*8)/8,
                              priorModelPrefixFits=(12*quantized<=math.floor(width*8)/8)))
for threshold in [14.6,15.1]:
    paired(f'kern-{threshold}-scaled20x72p5',[dict(run('AVATAR ToTo',20),kern=threshold)],width=290,wrap=False,fontScale=72.5,lineSpacingReduction=0)
for face in ['bold','serif']:
    paired(f'actual-{face}-fractional-width',[run('Agjp',14.5)],face=face,width=111.01,wrap=False)
for scale in [100,50]:
    paired(f'table-stored-scale{scale}',[run('Agjp',14.5)],table=True,width=111.01,wrap=False,fontScale=scale,lineSpacingReduction=0)
assert len(cases)==24

def properties(name,size):
    node=elem(name,sz=round(size*100),b='1' if selected_face['bold'] else '0',i='1' if selected_face['italic'] else '0',lang='en-US',kern='0')
    fill=elem('solidFill');fill.append(elem('srgbClr',val='000000'));node.append(fill)
    node.append(elem('latin',typeface=font_family))
    return node

deck=Presentation();deck.slide_width=Pt(720);deck.slide_height=Pt(720)
for index,case in enumerate(cases):
    selected_face=FACES[case.get('face','regular')];font_family=selected_face['family']
    if index%6==0:slide=deck.slides.add_slide(deck.slide_layouts[6])
    x=30+350*(index%2);y=50+220*((index%6)//2);w=case.get('width',290);h=160
    label=slide.shapes.add_textbox(Pt(x),Pt(y-23),Pt(310),Pt(18));label.text=case['name']
    label.text_frame.paragraphs[0].runs[0].font.size=Pt(10)
    if case.get('table'):
        shape=slide.shapes.add_table(1,1,Pt(x),Pt(y),Pt(w),Pt(h))
        cell=shape.table.cell(0,0);cell.margin_left=cell.margin_right=cell.margin_top=cell.margin_bottom=0
        tf=cell.text_frame;pr=shape.table._tbl.tblPr
        for child in list(pr):
            if child.tag=='{'+A+'}tableStyleId':pr.remove(child)
        pr.set('firstRow','0');pr.set('bandRow','0')
    else:
        shape=slide.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h));tf=shape.text_frame
    shape.name=case['name'];tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.margin_left=Pt(case.get('insets',[0,0,0,0])[0]);tf.margin_right=Pt(case.get('insets',[0,0,0,0])[1])
    tf.word_wrap=case.get('wrap',False);tf.auto_size=MSO_AUTO_SIZE.NONE
    body=tf._txBody;bodyPr=body.find('{'+A+'}bodyPr');bodyPr.set('anchor',case.get('anchor','t'))
    if 'compatLnSpc' in case:bodyPr.set('compatLnSpc',case['compatLnSpc'])
    if 'fontScale' in case:
        for c in list(bodyPr):
            if etree.QName(c).localname in ['noAutofit','normAutofit','spAutoFit']:bodyPr.remove(c)
        bodyPr.append(elem('normAutofit',fontScale=round(case['fontScale']*1000),lnSpcReduction=round(case['lineSpacingReduction']*1000)))
    for child in list(body):
        if child.tag in ['{'+A+'}p','{'+A+'}lstStyle']:body.remove(child)
    lst=elem('lstStyle');level=elem('lvl1pPr');level.append(properties('defRPr',case.get('listSize',18)));lst.append(level);body.append(lst)
    for specification in case['paragraphs']:
        p=elem('p');pr=elem('pPr',algn=case.get('alignment','l'),lvl=0)
        if 'tab' in case:
            tabs=elem('tabLst');tabs.append(elem('tab',pos=round(case['tab']['position']*12700),algn=case['tab']['alignment']));pr.append(tabs)
        spacing=specification.get('spacing')
        if spacing:
            line=elem('lnSpc')
            line.append(elem('spcPts',val=round(spacing['points']*100)) if 'points' in spacing else elem('spcPct',val=round(spacing['percent']*1000)))
            pr.append(line)
        if specification['default'] is not None:pr.append(properties('defRPr',specification['default']))
        p.append(pr)
        for token in specification['nodes']:
            if token['kind']=='break':
                node=elem('br')
                if token['size'] is not None:node.append(properties('rPr',token['size']))
            else:
                node=elem('r');rpr=properties('rPr',token['size'])
                if 'kern' in token:
                    if token['kern'] is None:rpr.attrib.pop('kern',None)
                    else:rpr.set('kern',str(round(token['kern']*100)))
                if 'tracking' in token:rpr.set('spc',str(round(token['tracking']*100)))
                if 'color' in token:rpr.find('{'+A+'}solidFill').find('{'+A+'}srgbClr').set('val',token['color'])
                node.append(rpr);text=elem('t');text.text=token['text'];node.append(text)
            p.append(node)
        if specification['end'] is not None:p.append(properties('endParaRPr',specification['end']))
        body.append(p)
    case.update(page=index//6,x=x,y=y,width=w,height=h)

# Independently embed actual licensed faces in their correct style slots.
buffer=BytesIO();deck.save(buffer)
with zipfile.ZipFile(buffer) as source:parts={n:source.read(n) for n in source.namelist()}
pres=etree.fromstring(parts['ppt/presentation.xml']);pres.set('embedTrueTypeFonts','1');pres.set('saveSubsetFonts','0')
fonts=etree.Element('{'+P+'}embeddedFontLst');entries={}
rels=etree.fromstring(parts['ppt/_rels/presentation.xml.rels']);face_manifest={}
for key,face in FACES.items():
    font_bytes=face['path'].read_bytes();font=TTFont(face['path']);header=bytearray(82)
    struct.pack_into('<IIII',header,0,0,len(font_bytes),0x00020001,0);header[26]=1;header[27]=1 if face['italic'] else 0
    struct.pack_into('<IHH',header,28,font['OS/2'].usWeightClass,font['OS/2'].fsType,0x504c)
    names=bytearray()
    for value in [face['family'],face['style'],font['name'].getDebugName(5),face['family']]:
        encoded=value.encode('utf-16le');names+=struct.pack('<H',len(encoded))+encoded+b'\0\0'
    eot=header+names+b'\0\0'+font_bytes;struct.pack_into('<I',eot,0,len(eot))
    if face['family'] not in entries:
        entry=etree.SubElement(fonts,'{'+P+'}embeddedFont');etree.SubElement(entry,'{'+P+'}font').set('typeface',face['family']);entries[face['family']]=entry
    relationship='rIdPaint'+key.title()
    etree.SubElement(entries[face['family']],'{'+P+'}'+face['slot']).set('{'+R+'}id',relationship)
    etree.SubElement(rels,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',Id=relationship,Type=R+'/font',Target='fonts/'+key+'.fntdata')
    parts['ppt/fonts/'+key+'.fntdata']=eot
    # Source faces are shared with the already committed NativeListMarkers fixture.
    face_manifest[key]=dict(family=face['family'],bold=face['bold'],italic=face['italic'],file='../NativeListMarkers/fonts/'+face['path'].name,sha256=hashlib.sha256(font_bytes).hexdigest(),embeddingPermissions=font['OS/2'].fsType)
pres.insert(list(pres).index(pres.find('{'+P+'}defaultTextStyle')),fonts)
def xml(node):return etree.tostring(node,xml_declaration=True,encoding='UTF-8',standalone=True)
parts['ppt/presentation.xml']=xml(pres);parts['ppt/_rels/presentation.xml.rels']=xml(rels)
ct=etree.fromstring(parts['[Content_Types].xml']);etree.SubElement(ct,'{http://schemas.openxmlformats.org/package/2006/content-types}Default',Extension='fntdata',ContentType='application/x-fontdata');parts['[Content_Types].xml']=xml(ct)
output=OUT/'native-body-alignment-v1.pptx'
with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as target:
    for name,data in sorted(parts.items()):target.writestr(zipfile.ZipInfo(name,date_time=(2026,10,4,0,0,0)),data,compress_type=zipfile.ZIP_DEFLATED)
(OUT/'cases.json').write_text(json.dumps(cases,indent=2)+'\n')
manifest=dict(source=output.name,sourceSHA256=hashlib.sha256(output.read_bytes()).hexdigest(),slideCount=4,caseCount=len(cases),faces=face_manifest,scope='Independent centered/right body offset, wrapping, padding, real-face/style, positive kerning threshold and table stored-scale calibration. No alignment formula or native outcome assumed.')
manifest['pdfSHA256']='49b3001cee2777ec1c319c191f40836c4c20391a65224b49f6289d54908b6b66'
(OUT/'manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
print(output);print(manifest['sourceSHA256']);print('4 slides, 24 cases')
