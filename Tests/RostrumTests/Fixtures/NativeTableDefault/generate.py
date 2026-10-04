"""Independent absent-versus-applied table-style native discriminators."""
from pathlib import Path
from io import BytesIO
import hashlib,json,struct,zipfile
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from fontTools.ttLib import TTFont
OUT=Path(__file__).resolve().parent
ROOT=next(p for p in OUT.parents if (p/'Package.swift').exists())
FONT=ROOT/'Tests/RostrumTests/Fixtures/NativeListMarkers/fonts/DejaVuSans.ttf'
A='http://schemas.openxmlformats.org/drawingml/2006/main';P='http://schemas.openxmlformats.org/presentationml/2006/main';R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
MEDIUM='{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}';GRID='{5940675A-B579-460E-94D1-54222C63F5DA}';NONE='{2D5ABB26-0587-4C30-8999-92F81FD0307C}';CUSTOM='{51477201-257E-4E39-9283-B9AAD00E8017}'
def element(name,**attrs):return etree.Element('{'+A+'}'+name,**{k:str(v) for k,v in attrs.items()})
def solid(color):
 n=element('solidFill');n.append(element('srgbClr',val=color));return n
def line(name,color,width):
 n=element(name,w=round(width*12700),cap='flat',cmpd='sng',algn='ctr');n.append(solid(color));return n
def custom(name='tblStyle'):
 n=element(name,styleId=CUSTOM,styleName='Owned purple cyan control');whole=element('wholeTbl');style=element('tcStyle');borders=element('tcBdr')
 for edge in ['left','right','top','bottom','insideH','insideV']:
  wrapper=element(edge);wrapper.append(line('ln','00AACC',2));borders.append(wrapper)
 style.append(borders);fill=element('fill');fill.append(solid('CC44AA'));style.append(fill);whole.append(style);n.append(whole);return n
fontdata=FONT.read_bytes();font=TTFont(FONT)
fontheader=bytearray(82);struct.pack_into('<IIII',fontheader,0,0,len(fontdata),0x00020001,0);fontheader[26]=1
struct.pack_into('<IHH',fontheader,28,font['OS/2'].usWeightClass,font['OS/2'].fsType,0x504c)
names=bytearray()
for value in ['DejaVu Sans','Book',font['name'].getDebugName(5),'DejaVu Sans']:
 b=value.encode('utf-16le');names+=struct.pack('<H',len(b))+b+b'\0\0'
eot=fontheader+names+b'\0\0'+fontdata;struct.pack_into('<I',eot,0,len(eot))
def xml(n):return etree.tostring(n,xml_declaration=True,encoding='UTF-8',standalone=True)
for group,default,choices in [('builtin',MEDIUM,[('absent',None),('explicit-medium',MEDIUM),('explicit-grid',GRID),('explicit-no-grid',NONE)]),('custom',CUSTOM,[('absent',None),('explicit-custom',CUSTOM),('absent-direct',None),('inline-custom','inline')])]:
 out=OUT/group;out.mkdir(exist_ok=True)
 deck=Presentation();deck.slide_width=Pt(720);deck.slide_height=Pt(720);slide=deck.slides.add_slide(deck.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string('DDEECC')
 cases=[]
 for i,(name,styleid) in enumerate(choices):
  x=30+350*(i%2);y=70+280*(i//2);width=111.01;height=160
  label=slide.shapes.add_textbox(Pt(x),Pt(y-28),Pt(310),Pt(20));label.text=group+'-'+name;label.text_frame.paragraphs[0].runs[0].font.size=Pt(12);label.text_frame.paragraphs[0].runs[0].font.name='DejaVu Sans'
  shape=slide.shapes.add_table(1,1,Pt(x),Pt(y),Pt(width),Pt(height));shape.name=group+'-'+name;table=shape.table;pr=table._tbl.tblPr
  for c in list(pr):pr.remove(c)
  pr.set('firstRow','0');pr.set('bandRow','0')
  if styleid=='inline':pr.append(custom('tableStyle'))
  elif styleid is not None:
   idnode=element('tableStyleId');idnode.text=styleid;pr.append(idnode)
  cell=table.cell(0,0);cell.margin_left=cell.margin_right=cell.margin_top=cell.margin_bottom=0
  body=cell.text_frame._txBody
  for child in list(body):body.remove(child)
  bp=element('bodyPr',anchor='t',lIns=0,rIns=0,tIns=0,bIns=0,wrap='none');bp.append(element('normAutofit',fontScale=100000,lnSpcReduction=0));body.append(bp);body.append(element('lstStyle'))
  paragraph=element('p');paragraph.append(element('pPr',algn='ctr'));run=element('r');rp=element('rPr',sz=1450,kern=0,b=0,i=0,lang='en-US');rp.append(solid('000000'));rp.append(element('latin',typeface='DejaVu Sans'));run.append(rp);t=element('t');t.text='Agjp';run.append(t);paragraph.append(run);body.append(paragraph)
  cp=cell._tc.get_or_add_tcPr()
  if name=='absent-direct':
   cp.append(line('lnL','008800',2));right=element('lnR');right.append(element('noFill'));cp.append(right);cp.append(solid('FFCC66'))
  cases.append(dict(id=shape.name,page=0,x=x,y=y,width=width,height=height,text='Agjp',font='DejaVu Sans',authoredSize=14.5,kern=0,alignment='ctr',styleID=styleid,directOverrides=name=='absent-direct',bodyXML=etree.tostring(body).decode(),tablePropertiesXML=etree.tostring(pr).decode(),cellPropertiesXML=etree.tostring(cp).decode()))
 buf=BytesIO();deck.save(buf)
 with zipfile.ZipFile(buf) as z:parts={n:z.read(n) for n in z.namelist()}
 styles=element('tblStyleLst',def_=default);styles.attrib.pop('def_');styles.set('def',default)
 if group=='custom':styles.append(custom())
 parts['ppt/tableStyles.xml']=xml(styles)
 pres=etree.fromstring(parts['ppt/presentation.xml']);pres.set('embedTrueTypeFonts','1');pres.set('saveSubsetFonts','0');fonts=etree.Element('{'+P+'}embeddedFontLst');entry=etree.SubElement(fonts,'{'+P+'}embeddedFont');etree.SubElement(entry,'{'+P+'}font',typeface='DejaVu Sans');etree.SubElement(entry,'{'+P+'}regular').set('{'+R+'}id','rIdMarkerFont');pres.insert(list(pres).index(pres.find('{'+P+'}defaultTextStyle')),fonts);parts['ppt/presentation.xml']=xml(pres)
 rel=etree.fromstring(parts['ppt/_rels/presentation.xml.rels']);etree.SubElement(rel,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',Id='rIdMarkerFont',Type=R+'/font',Target='fonts/regular.fntdata');parts['ppt/_rels/presentation.xml.rels']=xml(rel);parts['ppt/fonts/regular.fntdata']=eot
 ct=etree.fromstring(parts['[Content_Types].xml']);etree.SubElement(ct,'{http://schemas.openxmlformats.org/package/2006/content-types}Default',Extension='fntdata',ContentType='application/x-fontdata');parts['[Content_Types].xml']=xml(ct)
 source=out/('native-table-default-'+group+'-v1.pptx')
 with zipfile.ZipFile(source,'w',zipfile.ZIP_DEFLATED) as z:
  for name,data in sorted(parts.items()):z.writestr(zipfile.ZipInfo(name,date_time=(2026,10,4,0,0,0)),data,compress_type=zipfile.ZIP_DEFLATED)
 manifest=dict(source=source.name,sourceSHA256=hashlib.sha256(source.read_bytes()).hexdigest(),caseCount=4,pageCount=1,defaultStyle=default,background='DDEECC',fontSHA256=hashlib.sha256(fontdata).hexdigest(),fontFile='../../NativeListMarkers/fonts/DejaVuSans.ttf',scope='Independent style activation and cell paint controls. Native results not assumed.')
 manifest['pdfSHA256']={'builtin':'8b668d179dbf79a9d77265c24e986cab648244ab7ea134e89e4ebbedf7170b27','custom':'db76f7094ee5d2eea822aac578f1a21e30389bd73f9dc5ba51b09c0e2246ee31'}[group]
 (out/'cases.json').write_text(json.dumps(cases,indent=2)+'\n');(out/'manifest.json').write_text(json.dumps(manifest,indent=2)+'\n');print(source);print(manifest['sourceSHA256'])
 with zipfile.ZipFile(source) as z:
  assert z.testzip() is None
  for name in z.namelist():
   if name.endswith(('.xml','.rels')):etree.fromstring(z.read(name))
 reopened=Presentation(source);assert len(reopened.slides)==1 and len(reopened.slides[0].shapes)==8
