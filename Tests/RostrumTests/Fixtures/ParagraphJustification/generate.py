"""Independent python-pptx input; native PowerPoint export is the layout oracle."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement

out = Path(__file__).parent
p = Presentation()
p.slide_width, p.slide_height = Pt(720), Pt(540)
s = p.slides.add_slide(p.slide_layouts[6])
texts = [
    ('regular', 'Alpha beta gamma delta epsilon zeta eta theta iota kappa lambda mu.'),
    ('mixed', 'Alpha beta gamma delta epsilon zeta eta theta iota kappa lambda mu.'),
    ('hard-break', 'Alpha beta\vgamma delta epsilon zeta eta theta iota kappa.'),
    ('edge-spaces', '  Alpha  beta gamma delta epsilon zeta eta theta iota kappa.  '),
    ('one-word', 'SupercalifragilisticexpialidociousSupercalifragilisticexpialidocious'),
    ('bullet', 'Alpha beta gamma delta epsilon zeta eta theta iota kappa lambda mu.'),
    ('tab', 'Alpha\tbeta gamma delta epsilon zeta eta theta iota kappa.'),
    ('last-line', 'Alpha beta.'),
]
for i, (name, text) in enumerate(texts):
    x, y = 30 + (i % 2) * 350, 20 + (i // 2) * 125
    label = s.shapes.add_textbox(Pt(x), Pt(y), Pt(310), Pt(20))
    label.text = name
    label.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
    box = s.shapes.add_textbox(Pt(x), Pt(y+22), Pt(300), Pt(94))
    box.name = name
    tf = box.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.text = text
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.JUSTIFY
    para.space_before = para.space_after = Pt(0)
    para.font.name = 'Arial'
    para.font.size = Pt(18)
    if name == 'mixed':
        para.clear()
        for value, bold in [('Alpha beta ',False),('gamma delta ',True),('epsilon zeta eta theta iota kappa lambda mu.',False)]:
            run = para.add_run(); run.text=value; run.font.bold=bold
    if name == 'bullet':
        pp=para._p.get_or_add_pPr(); pp.set('marL', str(Pt(18))); pp.set('indent', str(Pt(-14)))
        bullet=OxmlElement('a:buChar'); bullet.set('char','•'); pp.insert(list(pp).index(pp.find('{http://schemas.openxmlformats.org/drawingml/2006/main}defRPr')), bullet)
p.save(out / 'paragraph-justification-v2.pptx')
