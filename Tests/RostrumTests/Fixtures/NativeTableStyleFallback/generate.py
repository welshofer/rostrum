"""Independent native partial custom-style missing-edge discriminators."""
from pathlib import Path
from io import BytesIO
import hashlib,json,struct,zipfile
from lxml import etree
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from fontTools.ttLib import TTFont
OUT=Path(__file__).resolve().parent
ROOT=next(p for p in OUT.parents if (p/'Package.swift').exists())
FONT=ROOT/'Tests/RostrumTests/Fixtures/NativeListMarkers/fonts/DejaVuSans.ttf'
A='http://schemas.openxmlformats.org/drawingml/2006/main';P='http://schemas.openxmlformats.org/presentationml/2006/main';R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
MEDIUM='{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}';GRID='{5940675A-B579-460E-94D1-54222C63F5DA}';NONE='{2D5ABB26-0587-4C30-8999-92F81FD0307C}';CUSTOM='{51477201-257E-4E39-9283-B9AAD00E8017}'
def element(name,**attrs):return etree.Element('{'+A+'}'+name,**{k:str(v) for k,v in attrs.items()})
def solid(color):
 n=element('solidFill');n.append(element('srgbClr',val=color));return n
def line(name,color,width):
 n=element(name,w=round(width*12700),cap='flat',cmpd='sng',algn='ctr');n.append(solid(color));return n
def custom(name='tblStyle'):
 n=element(name,styleId=CUSTOM,styleName='Owned purple cyan control');whole=element('wholeTbl');style=element('tcStyle');borders=element('tcBdr')
 for edge in ['left','right','top','bottom','insideH','insideV']:
  wrapper=element(edge);wrapper.append(line('ln','00AACC',2));borders.append(wrapper)
 style.append(borders);fill=element('fill');fill.append(solid('CC44AA'));style.append(fill);whole.append(style);n.append(whole);return n
fontdata=FONT.read_bytes();font=TTFont(FONT)
fontheader=bytearray(82);struct.pack_into('<IIII',fontheader,0,0,len(fontdata),0x00020001,0);fontheader[26]=1
struct.pack_into('<IHH',fontheader,28,font['OS/2'].usWeightClass,font['OS/2'].fsType,0x504c)
names=bytearray()
for value in ['DejaVu Sans','Book',font['name'].getDebugName(5),'DejaVu Sans']:
 b=value.encode('utf-16le');names+=struct.pack('<H',len(b))+b+b'\0\0'
eot=fontheader+names+b'\0\0'+fontdata;struct.pack_into('<I',eot,0,len(eot))
def xml(n):return etree.tostring(n,xml_declaration=True,encoding='UTF-8',standalone=True)

def make_properties(size=14.5):
 rp=element('rPr',sz=round(size*100),kern=0,b=0,i=0,lang='en-US');rp.append(solid('000000'));rp.append(element('latin',typeface='DejaVu Sans'));return rp


specs=[
 dict(id='absent-cell-style',kind='absent'),
 dict(id='empty-cell-style',kind='empty'),
 dict(id='empty-cell-borders',kind='empty-borders'),
 dict(id='referenced-fill-only',kind='fill'),
 dict(id='style-left-only',kind='left'),
 dict(id='style-empty-left-line',kind='empty-line'),
 dict(id='style-left-noFill',kind='noFill-line'),
 dict(id='inline-fill-only',kind='fill',inline=True),
 dict(id='fill-direct-left',kind='fill',direct='left'),
 dict(id='fill-direct-right-noFill',kind='fill',direct='noFill'),
 dict(id='partial-grid-whole',kind='partial-grid',rows=2,columns=2),
 dict(id='partial-grid-first-row-fill',kind='partial-grid',rows=2,columns=2,firstRow=True),
]
def fill_part(color):
 f=element('fill');f.append(solid(color));return f
def make_style(case,index):
 sid='{'+f'19000000-0000-4000-8000-{index+1:012d}'+'}'
 n=element('tblStyle',styleId=sid,styleName=case['id']);whole=element('wholeTbl');n.append(whole)
 kind=case['kind']
 if kind=='absent':return n
 style=element('tcStyle');whole.append(style)
 if kind=='empty':return n
 if kind!='fill':
  borders=element('tcBdr');style.append(borders)
  if kind in ['left','empty-line','noFill-line','partial-grid']:
   wrapper=element('left')
   if kind=='empty-line':ln=element('ln')
   elif kind=='noFill-line':
    ln=element('ln',w=25400);ln.append(element('noFill'))
   else:ln=line('ln','008800',2)
   wrapper.append(ln);borders.append(wrapper)
  if kind=='partial-grid':
   wrapper=element('insideH');wrapper.append(line('ln','0044CC',3));borders.append(wrapper)
 if kind!='empty-borders':style.append(fill_part('CC44AA'))
 if case.get('firstRow'):
  region=element('firstRow');region_style=element('tcStyle');region_style.append(fill_part('FFCC66'));region.append(region_style);n.append(region)
 return n
style_list=element('tblStyleLst');style_list.set('def',MEDIUM)
deck=Presentation();deck.slide_width=Pt(720);deck.slide_height=Pt(720)
for page in range(2):
 slide=deck.slides.add_slide(deck.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string('DDEECC')
 for i,case in enumerate(specs[page*6:page*6+6]):
  index=page*6+i;x=30+360*(i%2);y=65+215*(i//2);w=240;h=140
  definition=make_style(case,index);style_list.append(definition)
  label=slide.shapes.add_textbox(Pt(x),Pt(y-27),Pt(325),Pt(20));label.text=case['id'];label.text_frame.paragraphs[0].runs[0].font.size=Pt(12);label.text_frame.paragraphs[0].runs[0].font.name='DejaVu Sans'
  rows=case.get('rows',1);cols=case.get('columns',1)
  shape=slide.shapes.add_table(rows,cols,Pt(x),Pt(y),Pt(w),Pt(h));shape.name=case['id'];table=shape.table
  pr=table._tbl.tblPr
  for child in list(pr):pr.remove(child)
  pr.set('firstRow','1' if case.get('firstRow') else '0');pr.set('bandRow','0')
  if case.get('inline'):
   inline=etree.fromstring(etree.tostring(definition));inline.tag='{'+A+'}tableStyle';pr.append(inline)
  else:
   choice=element('tableStyleId');choice.text=definition.get('styleId');pr.append(choice)
  cells=[]
  for row in range(rows):
   for col in range(cols):
    cell=table.cell(row,col);cell.margin_left=cell.margin_right=cell.margin_top=cell.margin_bottom=0
    cp=cell._tc.get_or_add_tcPr()
    for child in list(cp):cp.remove(child)
    if case.get('direct')=='left':cp.append(line('lnL','0044CC',4))
    elif case.get('direct')=='noFill':
     ln=element('lnR');ln.append(element('noFill'));cp.append(ln)
    body=cell.text_frame._txBody
    for child in list(body):body.remove(child)
    bp=element('bodyPr',anchor='t',lIns=0,rIns=0,tIns=0,bIns=0,wrap='none');bp.append(element('normAutofit',fontScale=100000,lnSpcReduction=0));body.append(bp);body.append(element('lstStyle'))
    paragraph=element('p');paragraph.append(element('pPr',algn='ctr'));run=element('r');run.append(make_properties());t=element('t');t.text='Agjp';run.append(t);paragraph.append(run);body.append(paragraph)
    cells.append(dict(row=row,column=col,text='Agjp',properties=etree.tostring(cp).decode(),textBody=etree.tostring(body).decode()))
  case.update(page=page,x=x,y=y,width=w,height=h,rows=rows,columns=cols,cells=cells,styleDefinition=etree.tostring(definition).decode(),tableProperties=etree.tostring(pr).decode())
buf=BytesIO();deck.save(buf)
with zipfile.ZipFile(buf) as z:parts={n:z.read(n) for n in z.namelist()}
parts['ppt/tableStyles.xml']=xml(style_list)
pres=etree.fromstring(parts['ppt/presentation.xml']);pres.set('embedTrueTypeFonts','1');pres.set('saveSubsetFonts','0');fonts=etree.Element('{'+P+'}embeddedFontLst');entry=etree.SubElement(fonts,'{'+P+'}embeddedFont');etree.SubElement(entry,'{'+P+'}font',typeface='DejaVu Sans');etree.SubElement(entry,'{'+P+'}regular').set('{'+R+'}id','rIdMarkerFont');pres.insert(list(pres).index(pres.find('{'+P+'}defaultTextStyle')),fonts);parts['ppt/presentation.xml']=xml(pres)
rel=etree.fromstring(parts['ppt/_rels/presentation.xml.rels']);etree.SubElement(rel,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',Id='rIdMarkerFont',Type=R+'/font',Target='fonts/regular.fntdata');parts['ppt/_rels/presentation.xml.rels']=xml(rel);parts['ppt/fonts/regular.fntdata']=eot
ct=etree.fromstring(parts['[Content_Types].xml']);etree.SubElement(ct,'{http://schemas.openxmlformats.org/package/2006/content-types}Default',Extension='fntdata',ContentType='application/x-fontdata');parts['[Content_Types].xml']=xml(ct)
output=OUT/'native-table-style-fallback19-v1.pptx'
with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as z:
 for name,data in sorted(parts.items()):z.writestr(zipfile.ZipInfo(name,date_time=(2026,10,4,0,0,0)),data,compress_type=zipfile.ZIP_DEFLATED)
manifest=dict(expectedVisibleGlyphs=sum(len(cell['text']) for case in specs for cell in case['cells']),source=output.name,sourceSHA256=hashlib.sha256(output.read_bytes()).hexdigest(),cases=12,slides=2,fontSHA256=hashlib.sha256(fontdata).hexdigest(),scope='Research only: missing custom-style borders versus empty/explicit noFill declarations; no fallback rule assumed.')
(OUT/'manifest.json').write_text(json.dumps(manifest,indent=2)+'\n');(OUT/'cases.json').write_text(json.dumps(specs,indent=2)+'\n')
with zipfile.ZipFile(output) as z:
 assert z.testzip() is None
 for name in z.namelist():
  if name.endswith(('.xml','.rels')):etree.fromstring(z.read(name))
assert len(Presentation(output).slides)==2
print(output);print(manifest['sourceSHA256'])
