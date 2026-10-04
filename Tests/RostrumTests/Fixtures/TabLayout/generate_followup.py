"""Independent tab probes; open in native PowerPoint and export PDF."""
from pathlib import Path
import json
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
out=Path(__file__).parent
p=Presentation(); p.slide_width=Pt(720); p.slide_height=Pt(540)
cases=[
 ('center-short','A\tBeta gamma delta',[(220,'ctr')],{}),
 ('center-overflow','A\tBeta gamma delta epsilon',[(220,'ctr')],{}),
 ('center-long','A\tBeta gamma delta epsilon zeta eta theta iota kappa',[(220,'ctr')],{}),
 ('center-long-nowrap','A\tBeta gamma delta epsilon zeta eta theta iota kappa',[(220,'ctr')],{'nowrap':True}),
 ('center-middle','A\tBeta gamma delta epsilon zeta eta theta iota kappa',[(150,'ctr')],{}),
 ('center-middle-nowrap','A\tBeta gamma delta epsilon zeta eta theta iota kappa',[(150,'ctr')],{'nowrap':True}),
 ('center-prefix','Alpha beta\tBeta gamma delta epsilon',[(220,'ctr')],{}),
 ('center-overflow-nowrap','A\tBeta gamma delta epsilon',[(220,'ctr')],{'nowrap':True}),
 ('center-hard','A\tBeta gamma delta epsilon\vE F',[(220,'ctr')],{}),
 ('center-boundary','A\tBeta gamma delta epsilon zeta',[(180,'ctr')],{}),
 ('center-tabnext','A\tBeta gamma delta epsilon\tZ',[(180,'ctr'),(290,'l')],{}),
 ('center-kerning','A\tAVATAR office',[(180,'ctr')],{'mixed':False}),
]
for part in p.part.package.iter_parts():
 if hasattr(part, '_element'):
  for element in part._element.iter(): element.attrib.pop('defTabSz',None)
records=[]
for i,(name,text,tabs,opts) in enumerate(cases):
 if i%6==0: s=p.slides.add_slide(p.slide_layouts[6])
 x,y=30+(i%2)*350,20+((i%6)//2)*170
 label=s.shapes.add_textbox(Pt(x),Pt(y),Pt(300),Pt(18)); label.text=name; label.text_frame.paragraphs[0].runs[0].font.size=Pt(10)
 box=s.shapes.add_textbox(Pt(x),Pt(y+25),Pt(300),Pt(125)); box.name=name
 tf=box.text_frame; tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 tf.word_wrap=not opts.get('nowrap'); tf.auto_size=MSO_AUTO_SIZE.NONE; tf.text=text
 para=tf.paragraphs[0]; para.font.name='Arial'; para.font.size=Pt(18); para.space_before=para.space_after=Pt(0)
 if opts.get('justify'): para.alignment=PP_ALIGN.JUSTIFY
 pp=para._p.get_or_add_pPr()
 for key in ['marL','indent']:
  if key in opts: pp.set(key,str(Pt(opts[key])))
 if 'interval' in opts: pp.set('defTabSz',str(Pt(opts['interval'])))
 if opts.get('bullet'):
  b=OxmlElement('a:buChar'); b.set('char','•'); pp.insert(2,b)
 if opts.get('inherited'):
  style=tf._txBody.find('{http://schemas.openxmlformats.org/drawingml/2006/main}lstStyle')
  level=OxmlElement('a:lvl1pPr'); level.set('defTabSz',str(Pt(48)))
  lst=OxmlElement('a:tabLst'); stop=OxmlElement('a:tab'); stop.set('pos',str(Pt(100))); stop.set('algn','r'); lst.append(stop); level.append(lst); style.append(level)
 if tabs or opts.get('empty'):
  tl=OxmlElement('a:tabLst')
  for pos,align in tabs:
   t=OxmlElement('a:tab'); t.set('pos',str(Pt(pos))); t.set('algn',align); tl.append(t)
  pp.insert(list(pp).index(pp.find('{http://schemas.openxmlformats.org/drawingml/2006/main}defRPr')),tl)
 if opts.get('mixed'):
  para.clear()
  for val,bold in [('A\t',False),('Beta',True),(' Z',False)]:
   r=para.add_run(); r.text=val; r.font.bold=bold
 records.append(dict(name=name,text=text,tabs=tabs,options=opts,page=i//6,x=x,y=y+25,width=300,height=125))
p.save(out/'tab-layout-v3.pptx')
(out/'cases-v3.json').write_text(json.dumps(records,indent=2)+'\n')
