"""Independent tab probes; open in native PowerPoint and export PDF."""
from pathlib import Path
import json
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
out=Path(__file__).parent
p=Presentation(); p.slide_width=Pt(720); p.slide_height=Pt(540)
cases=[
 ('default','A\tB\tC',[],{}),
 ('left','A\tBeta',[(120,'l')],{}),
 ('center','A\tBeta',[(120,'ctr')],{}),
 ('right','A\tBeta',[(120,'r')],{}),
 ('decimal','A\t123.45',[(120,'dec')],{}),
 ('decimal-integer','A\t123',[(120,'dec')],{}),
 ('multiple','A\tBeta\t123.45',[(100,'ctr'),(210,'dec')],{}),
 ('mixed','A\tBeta Z',[(120,'r')],{'mixed':True}),
 ('indent','A\tBeta',[(120,'l')],{'marL':24,'indent':12}),
 ('bullet','A\tBeta',[(120,'r')],{'marL':24,'indent':-18,'bullet':True}),
 ('default-interval','A\tB\tC',[],{'interval':48}),
 ('custom-then-default','A\tB\tC\tD',[(50,'l')],{'interval':48}),
 ('justify-left','Alpha beta\tgamma delta epsilon zeta eta theta iota kappa lambda mu.',[(120,'l')],{'justify':True}),
 ('justify-right','Alpha beta\tgamma delta epsilon zeta eta theta iota kappa lambda mu.',[(180,'r')],{'justify':True}),
 ('justify-multiple','A B\tC D\tE F G H I J K L M N O P Q R S T U V W X Y Z',[(70,'l'),(140,'l')],{'justify':True}),
 ('wrap-left','A\tBeta gamma delta epsilon zeta eta theta iota kappa',[(220,'l')],{}),
 ('wrap-center','A\tBeta gamma delta epsilon zeta eta theta iota kappa',[(220,'ctr')],{}),
 ('wrap-right','A\tBeta gamma delta epsilon zeta eta theta iota kappa',[(220,'r')],{}),
 ('collision-right','Alpha beta gamma\tBeta',[(120,'r')],{}),
 ('hard-break','A B\tC D\vE F',[(120,'l')],{'justify':True}),
 ('decimal-comma','A\t123,45',[(120,'dec')],{}),
 ('decimal-suffix','A\t123.45 USD',[(120,'dec')],{}),
 ('trailing-spaces','A\tBeta  ',[(120,'r')],{}),
 ('nowrap','A\tBeta gamma delta epsilon zeta eta theta iota kappa',[(220,'r')],{'nowrap':True}),
]
# Remove every template-provided default so the bare-default probe isolates the
# application fallback, not python-pptx's explicit inherited 36pt interval.
for part in p.part.package.iter_parts():
 if hasattr(part, '_element'):
  for element in part._element.iter(): element.attrib.pop('defTabSz',None)
cases += [
 ('inherited-tabs','A\tBeta',[],{'inherited':True}),
 ('empty-tabs','A\tBeta',[],{'inherited':True,'empty':True}),
 ('override-tabs','A\tBeta',[(120,'l')],{'inherited':True}),
 ('beyond-width','A\tBeta',[(330,'l')],{}),
 ('center-trailing','A\tBeta  ',[(120,'ctr')],{}),
 ('decimal-period-leading','A\t.45',[(120,'dec')],{}),
]
records=[]
for i,(name,text,tabs,opts) in enumerate(cases):
 if i%6==0: s=p.slides.add_slide(p.slide_layouts[6])
 x,y=30+(i%2)*350,20+((i%6)//2)*170
 label=s.shapes.add_textbox(Pt(x),Pt(y),Pt(300),Pt(18)); label.text=name; label.text_frame.paragraphs[0].runs[0].font.size=Pt(10)
 box=s.shapes.add_textbox(Pt(x),Pt(y+25),Pt(300),Pt(125)); box.name=name
 tf=box.text_frame; tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 tf.word_wrap=not opts.get('nowrap'); tf.auto_size=MSO_AUTO_SIZE.NONE; tf.text=text
 para=tf.paragraphs[0]; para.font.name='Arial'; para.font.size=Pt(18); para.space_before=para.space_after=Pt(0)
 if opts.get('justify'): para.alignment=PP_ALIGN.JUSTIFY
 pp=para._p.get_or_add_pPr()
 for key in ['marL','indent']:
  if key in opts: pp.set(key,str(Pt(opts[key])))
 if 'interval' in opts: pp.set('defTabSz',str(Pt(opts['interval'])))
 if opts.get('bullet'):
  b=OxmlElement('a:buChar'); b.set('char','•'); pp.insert(2,b)
 if opts.get('inherited'):
  style=tf._txBody.find('{http://schemas.openxmlformats.org/drawingml/2006/main}lstStyle')
  level=OxmlElement('a:lvl1pPr'); level.set('defTabSz',str(Pt(48)))
  lst=OxmlElement('a:tabLst'); stop=OxmlElement('a:tab'); stop.set('pos',str(Pt(100))); stop.set('algn','r'); lst.append(stop); level.append(lst); style.append(level)
 if tabs or opts.get('empty'):
  tl=OxmlElement('a:tabLst')
  for pos,align in tabs:
   t=OxmlElement('a:tab'); t.set('pos',str(Pt(pos))); t.set('algn',align); tl.append(t)
  pp.insert(list(pp).index(pp.find('{http://schemas.openxmlformats.org/drawingml/2006/main}defRPr')),tl)
 if opts.get('mixed'):
  para.clear()
  for val,bold in [('A\t',False),('Beta',True),(' Z',False)]:
   r=para.add_run(); r.text=val; r.font.bold=bold
 records.append(dict(name=name,text=text,tabs=tabs,options=opts,page=i//6,x=x,y=y+25,width=300,height=125))
p.save(out/'tab-layout-v2.pptx')
(out/'cases-v2.json').write_text(json.dumps(records,indent=2)+'\n')
