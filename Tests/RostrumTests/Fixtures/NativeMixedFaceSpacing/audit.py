"""Audit the retained source/PDF pins and every owned frame/run property."""
from pathlib import Path
from io import BytesIO
import hashlib,json,zipfile,struct
from lxml import etree
from pptx import Presentation
from fontTools.ttLib import TTFont
R=Path(__file__).resolve().parent
A='{http://schemas.openxmlformats.org/drawingml/2006/main}'
receipts=[]
for folder,count in [('base',4),('anchors',8)]:
 r=R/folder;m=json.loads((r/'manifest.json').read_text());cases=json.loads((r/'cases.json').read_text());pptx=r/m['source'];d=Presentation(pptx)
 assert hashlib.sha256(pptx.read_bytes()).hexdigest()==m['sourceSHA256']
 assert hashlib.sha256((r/'powerpoint.pdf').read_bytes()).hexdigest()==m['pdfSHA256']
 assert len(d.slides)==1 and len(cases)==count and len({c['name'] for c in cases})==count
 with zipfile.ZipFile(pptx) as z:
  assert z.testzip() is None
  for name in z.namelist():
   if name.endswith(('.xml','.rels')):etree.fromstring(z.read(name))
  for key,face in m['faces'].items():
   data=(r/face['file']).read_bytes();assert hashlib.sha256(data).hexdigest()==face['sha256']
   eot=z.read('ppt/fonts/'+key+'.fntdata');size=struct.unpack_from('<I',eot,4)[0];assert eot[-size:]==data
   assert TTFont(BytesIO(data))['OS/2'].fsType==0
 for case in cases:
  shapes=[s for s in d.slides[0].shapes if s.name==case['name']];assert len(shapes)==1;shape=shapes[0]
  assert [shape.left,shape.top,shape.width,shape.height]==[round(case[k]*12700) for k in ['x','y','width','height']]
  body=shape.text_frame._txBody;bp=body.find(A+'bodyPr');assert bp.get('anchor')==case.get('anchor','t')
  assert all(bp.get(k)=='0' for k in ['lIns','tIns','rIns','bIns'])
  if 'fontScale' in case:
   auto=bp.find(A+'normAutofit');assert auto.get('fontScale')==str(round(case['fontScale']*1000));assert auto.get('lnSpcReduction')=='0'
  else:assert bp.find(A+'noAutofit') is not None
  assert body.find(A+'p').find(A+'pPr').find(A+'lnSpc').find(A+'spcPts').get('val')==str(round(case['paragraphs'][0]['spacing']['points']*100))
  nodes=body.find(A+'p');actual=[n for n in nodes if n.tag in [A+'r',A+'br']];expected=case['paragraphs'][0]['nodes'];assert len(actual)==len(expected)
  for node,run in zip(actual,expected):
   assert node.tag==A+('r' if run['kind']=='run' else 'br');pr=node.find(A+'rPr')
   assert pr.get('sz')==str(round(run['size']*100)) and pr.get('kern')=='0'
   face=m['faces'][run.get('face',case.get('face','regular'))];assert pr.find(A+'latin').get('typeface')==face['family']
   if run['kind']=='run':assert node.find(A+'t').text==run['text']
 receipts.append(dict(folder=folder,cases=count,sourceSHA256=m['sourceSHA256'],pdfSHA256=m['pdfSHA256']))
print(json.dumps(dict(status='passed',cases=12,checks='all frames, anchors, zero insets, autofit choices, exact spacing, node order, per-run faces/sizes/kerning/text, sfnt bytes and source/PDF pins',groups=receipts),indent=2))
