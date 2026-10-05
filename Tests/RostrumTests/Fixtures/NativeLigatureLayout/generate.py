"""Independent python-pptx/OOXML probes; no Rostrum measurements or writer."""
from pathlib import Path
from io import BytesIO
import hashlib,json,math,struct,zipfile
from lxml import etree
from fontTools.ttLib import TTFont
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import MSO_AUTO_SIZE,PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement

out=Path(__file__).parent
font_path=out.parent/'Typography/DejaVuSans.ttf'
font_bytes=font_path.read_bytes();font=TTFont(font_path)
cmap=font.getBestCmap();upem=font['head'].unitsPerEm

def run(text,size=18,kern=None,tracking=None,color='000000'):
 return dict(text=text,size=size,kern=kern,tracking=tracking,color=color)
def width(text,size):
 return sum(math.floor(font['hmtx'].metrics[cmap[ord(c)]][0]*size/upem*8+.5)/8 for c in text)
probes=[]
def add(name,runs,w=290,**options):probes.append(dict(name=name,runs=runs,width=w,**options))
for text in ['fi','fl','ff','ffi','ffl','office final waffle']:
 add('default-'+text.replace(' ','-'),[run(text)])
add('office-12-reference',[run('office',12,1200)])
add('kern-zero',[run('office final waffle',kern=0)])
add('tracking-positive',[run('office final waffle',tracking=.3)])
add('tracking-negative',[run('office final waffle',tracking=-.2)])
add('color-split',[run('of',color='AA3300'),run('fice',color='0033AA')])
add('same-style-split',[run('of'),run('fice')])
for suffix,delta in [('below',-.01),('above',.01)]:
 add('office-edge-'+suffix,[run('officeZ',kern=0)],width('office',18)+delta)
 add('ffi-component-edge-'+suffix,[run('ffiZ',kern=0)],width('ff',18)+delta)
 add('scaled-office-edge-'+suffix,[run('officeZ',kern=0)],width('office',9)+delta,fontScale=50000)
add('mixed-size',[run('of',18,0,color='AA3300'),run('fice',12,0,color='0033AA')])
add('mixed-size-edge',[run('of',18,0),run('ficeZ',12,0)],width('of',18)+width('fice',12)+.01)
add('table-default',[run('office final waffle')],120,table=True)
add('table-edge',[run('officeZ',kern=0)],width('office',18)+.01,table=True)
add('tab-field',[run('Label\toffice final waffle')],290,tab=90)
add('hard-break',[run('office\nfinal waffle')],150)

p=Presentation();p.slide_width=Pt(720);p.slide_height=Pt(540)
cases=[]
for i,probe in enumerate(probes):
 if i%6==0:slide=p.slides.add_slide(p.slide_layouts[6])
 x,y=30+(i%2)*350,20+((i%6)//2)*170
 label=slide.shapes.add_textbox(Pt(x),Pt(y),Pt(320),Pt(20));label.text=probe['name']
 label.text_frame.paragraphs[0].runs[0].font.size=Pt(10)
 y+=25;w=probe['width'];h=120
 if probe.get('table'):
  shape=slide.shapes.add_table(1,1,Pt(x),Pt(y),Pt(w),Pt(h));cell=shape.table.cell(0,0)
  cell.margin_left=cell.margin_right=cell.margin_top=cell.margin_bottom=0
  tf=cell.text_frame
  tp=shape.table._tbl.tblPr
  for child in list(tp):
   if child.tag.endswith('tableStyleId'):tp.remove(child)
  tp.set('firstRow','0');tp.set('bandRow','0')
 else:shape=slide.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h));tf=shape.text_frame
 shape.name=probe['name'];tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 tf.word_wrap=True;tf.auto_size=MSO_AUTO_SIZE.NONE
 para=tf.paragraphs[0];para.alignment=PP_ALIGN.LEFT;para.space_before=para.space_after=Pt(0)
 pp=para._p.get_or_add_pPr()
 if probe.get('tab'):
  tl=OxmlElement('a:tabLst');t=OxmlElement('a:tab');t.set('pos',str(Pt(probe['tab'])));t.set('algn','l');tl.append(t);pp.append(tl)
 if probe.get('fontScale'):
  bp=tf._txBody.find('{http://schemas.openxmlformats.org/drawingml/2006/main}bodyPr')
  for e in list(bp):
   if e.tag.endswith('noAutofit'):bp.remove(e)
  af=OxmlElement('a:normAutofit');af.set('fontScale',str(probe['fontScale']));bp.append(af)
 for r in probe['runs']:
  for part_index,part in enumerate(r['text'].split('\n')):
   if part_index:para.add_line_break()
   rr=para.add_run();rr.text=part;rr.font.name='DejaVu Sans';rr.font.size=Pt(r['size']);rr.font.color.rgb=RGBColor.from_string(r['color'])
   rp=rr._r.get_or_add_rPr();rp.set('lang','en-US')
   if probe.get('table'):rp.set('b','0');rp.set('i','0')
   if r['kern'] is not None:rp.set('kern',str(r['kern']))
   if r['tracking'] is not None:rp.set('spc',str(round(r['tracking']*100)))
 cases.append(dict(probe,page=i//6,x=x,y=y,width=shape.width/12700,height=h))
buf=BytesIO();p.save(buf)
# EOT v2.1 is independently constructed per W3C specification.
header=bytearray(82);struct.pack_into('<IIII',header,0,0,len(font_bytes),0x00020001,0);header[26]=1
struct.pack_into('<IHH',header,28,400,font['OS/2'].fsType,0x504c)
names=bytearray()
for value in ['DejaVu Sans','Book','Version 2.37','DejaVu Sans']:
 encoded=value.encode('utf-16le');names+=struct.pack('<H',len(encoded))+encoded+b'\0\0'
eot=header+names+b'\0\0'+font_bytes;struct.pack_into('<I',eot,0,len(eot))
NSP='http://schemas.openxmlformats.org/presentationml/2006/main';NSR='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
with zipfile.ZipFile(buf) as source:parts={n:source.read(n) for n in source.namelist()}
pres=etree.fromstring(parts['ppt/presentation.xml']);pres.set('embedTrueTypeFonts','1');pres.set('saveSubsetFonts','0')
el=etree.Element('{'+NSP+'}embeddedFontLst');ef=etree.SubElement(el,'{'+NSP+'}embeddedFont')
etree.SubElement(ef,'{'+NSP+'}font').set('typeface','DejaVu Sans');etree.SubElement(ef,'{'+NSP+'}regular').set('{'+NSR+'}id','rIdLigatureFont')
default=pres.find('{'+NSP+'}defaultTextStyle');pres.insert(list(pres).index(default),el)
parts['ppt/presentation.xml']=etree.tostring(pres,xml_declaration=True,encoding='UTF-8',standalone=True)
rels=etree.fromstring(parts['ppt/_rels/presentation.xml.rels']);etree.SubElement(rels,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',Id='rIdLigatureFont',Type=NSR+'/font',Target='fonts/ligatures.fntdata')
parts['ppt/_rels/presentation.xml.rels']=etree.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
ct=etree.fromstring(parts['[Content_Types].xml']);etree.SubElement(ct,'{http://schemas.openxmlformats.org/package/2006/content-types}Default',Extension='fntdata',ContentType='application/x-fontdata')
parts['[Content_Types].xml']=etree.tostring(ct,xml_declaration=True,encoding='UTF-8',standalone=True);parts['ppt/fonts/ligatures.fntdata']=eot
path=out/'native-ligatures-regular-table.pptx'
with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as target:
 for name,data in sorted(parts.items()):target.writestr(name,data)
(out/'cases.json').write_text(json.dumps(cases,indent=2)+'\n')
(out/'input-manifest.json').write_text(json.dumps(dict(generator='python-pptx 1.0.2 + fontTools; independent OOXML/EOT',source=path.name,sourceSHA256=hashlib.sha256(path.read_bytes()).hexdigest(),font=str(font_path),fontSHA256=hashlib.sha256(font_bytes).hexdigest(),slideCount=len(p.slides),caseCount=len(cases),unitsPerEm=upem,advances={chr(c):font['hmtx'].metrics[cmap[c]][0] for c in range(32,127)}),indent=2)+'\n')
print(path);print(len(p.slides),'slides',len(cases),'cases');print(hashlib.sha256(path.read_bytes()).hexdigest())
